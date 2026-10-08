from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from modules.comparador import ordenar_competencias
from modules.maternidade import resumir_maternidade, separar_periodos_sem_dados


AZUL = "1F4E78"
AZUL_CLARO = "D9EAF7"
MOEDA = '"R$" #,##0.00;[Red]("R$" #,##0.00);"R$" 0.00'


def _comparativo_maternidade(writer, comparativo: pd.DataFrame, parcial: bool) -> None:
    base, sem_dados = separar_periodos_sem_dados(_ordenar(comparativo))
    empresas = comparativo.cnpj.dropna().astype(str).unique()
    unica = len(empresas) == 1
    if unica:
        base = base.drop(columns="cnpj")
        sem_dados = sem_dados.drop(columns="cnpj")
    nomes = {"cnpj": "CNPJ", "competencia": "Competência", "Apoios considerados": "Apoios"}
    base = base.rename(columns=nomes)
    sem_dados = sem_dados.rename(columns=nomes)
    base.to_excel(writer, sheet_name="Comparativo", startrow=11, index=False)
    ws = writer.book["Comparativo"]
    ws["A1"] = "Conciliação de salário-maternidade"
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color=AZUL)
    ws["A2"] = f"CNPJ: {empresas[0]}" if unica else f"{len(empresas)} CNPJs; identificação mantida no comparativo"
    for linha, (rotulo, valor) in enumerate(resumir_maternidade(comparativo).items(), 4):
        ws.cell(linha, 1, rotulo)
        ws.cell(linha, 3, None if pd.isna(valor) else valor)
        ws.cell(linha, 3).font = Font(bold=True)
        ws.cell(linha, 3).number_format = MOEDA if linha < 6 else "0"
    for celula in ws[5][:3]:
        celula.fill = PatternFill("solid", fgColor="FFF2CC")
    ws["A9"] = "Diferença = total identificado − e-CAC. Positiva: potencial crédito; negativa: declarado superior. Valores em R$."
    ws["A10"] = "Conferência parcial; consultar ocorrências. " if parcial else "Conferência dos documentos recebidos. "
    ws["A10"] = ws["A10"].value + "AAAA indica 13º. Repetições preservadas precisam de revisão; apoios sujeitos à validação jurídica."
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B13" if not unica else "A13"
    fim = 12 + len(base)
    if len(base):
        ws.auto_filter.ref = f"A12:{get_column_letter(len(base.columns))}{fim}"
    for indice, nome in enumerate(base.columns, 1):
        coluna = get_column_letter(indice)
        ws.column_dimensions[coluna].width = 24 if nome not in ["CNPJ", "Pendência de revisão"] else (22 if nome == "CNPJ" else 56)
        monetaria = nome in ["Principal", "Apoios", "Total identificado", "Declarado no e-CAC", "Diferença"]
        for linha in range(13, fim + 1):
            item = ws.cell(linha, indice)
            item.font = Font(name="Arial", size=10)
            item.alignment = Alignment(vertical="center", wrap_text=not monetaria, horizontal="right" if monetaria else "left")
            if monetaria:
                item.number_format = MOEDA
            if linha % 2:
                item.fill = PatternFill("solid", fgColor=AZUL_CLARO)
            if nome == "Pendência de revisão" and item.value:
                item.fill = PatternFill("solid", fgColor="FFF2CC")
    ws.column_dimensions["A"].width = 42  # Rótulos do resumo, sem repetir CNPJ nas linhas.
    for linha in range(13, fim + 1):
        pendencia = str(ws.cell(linha, len(base.columns)).value or "")
        ws.row_dimensions[linha].height = max(32, ((len(pendencia) // 48) + 1) * 15 + 8)
    cabecalhos = [12]
    if not sem_dados.empty:
        secao = fim + 3
        ws.cell(secao, 1, "Períodos sem dados para comparação").font = Font(bold=True, color=AZUL)
        sem_dados.to_excel(writer, sheet_name="Comparativo", startrow=secao, index=False)
        cabecalhos.append(secao + 1)
        for linha in range(secao + 2, secao + 2 + len(sem_dados)):
            ws.row_dimensions[linha].height = 32
            for item in ws[linha][:len(sem_dados.columns)]:
                item.alignment = Alignment(vertical="center", wrap_text=True)
    for linha in cabecalhos:
        ws.row_dimensions[linha].height = 36
        quantidade = len(base.columns) if linha == 12 else len(sem_dados.columns)
        for item in ws[linha][:quantidade]:
            item.fill = PatternFill("solid", fgColor=AZUL)
            item.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            item.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ordenar(dados: pd.DataFrame) -> pd.DataFrame:
    if {"cnpj", "competencia"}.issubset(dados.columns):
        return ordenar_competencias(dados).reset_index(drop=True)
    return dados


def _ajustar_planilha(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for celula in ws[1]:
        celula.fill = PatternFill("solid", fgColor=AZUL)
        celula.font = Font(color="FFFFFF", bold=True)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 36
    for coluna in ws.columns:
        valores = [str(c.value or "") for c in coluna[:200]]
        largura = min(max(max(map(len, valores), default=0) + 2, 12), 48)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = largura


def gerar_excel_ecac(
    resumo: pd.DataFrame,
    documentos: pd.DataFrame,
    ocorrencias: pd.DataFrame,
) -> bytes:
    memoria = io.BytesIO()
    with pd.ExcelWriter(memoria, engine="openpyxl") as writer:
        aba = "Salário Maternidade" if "salario_maternidade" in resumo else "Salário Família"
        exibicao = resumo.drop(columns="salario_familia", errors="ignore") if "salario_maternidade" in resumo else resumo
        _ordenar(exibicao).to_excel(writer, sheet_name=aba, index=False)
        _ordenar(documentos).to_excel(writer, sheet_name="Documentos", index=False)
        if not ocorrencias.empty:
            ocorrencias.to_excel(writer, sheet_name="Ocorrências", index=False)
        for ws in writer.book.worksheets:
            _ajustar_planilha(ws)
        principal = writer.book[aba]
        cabecalhos = {celula.value: celula.column for celula in principal[1]}
        coluna_valor = cabecalhos.get("salario_maternidade", cabecalhos.get("salario_familia"))
        if coluna_valor:
            for celula in principal.iter_cols(
                min_col=coluna_valor,
                max_col=coluna_valor,
                min_row=2,
            ):
                for item in celula:
                    item.number_format = 'R$ #,##0.00'
            principal.column_dimensions[get_column_letter(coluna_valor)].width = 20
    return memoria.getvalue()


def gerar_excel_comparativo(
    comparativo: pd.DataFrame,
    ecac: pd.DataFrame,
    levantamento: pd.DataFrame,
    detalhe: pd.DataFrame | None = None,
    classificacao: pd.DataFrame | None = None,
    ocorrencias: pd.DataFrame | None = None,
) -> bytes:
    memoria = io.BytesIO()
    with pd.ExcelWriter(memoria, engine="openpyxl") as writer:
        maternidade = {"Principal", "Apoios considerados", "Diferença", "Pendência de revisão"}.issubset(comparativo.columns)
        if maternidade:
            parcial = (ocorrencias is not None and not ocorrencias.empty) or comparativo["Pendência de revisão"].str.contains("repeti|identificação parcial|divergente|inválido|conferência", case=False, na=False).any()
            _comparativo_maternidade(writer, comparativo, bool(parcial))
        else:
            _ordenar(comparativo).to_excel(writer, sheet_name="Comparativo", index=False)
        _ordenar(ecac).to_excel(writer, sheet_name="Dados e-CAC", index=False)
        _ordenar(levantamento).to_excel(writer, sheet_name="Levantamento", index=False)
        if detalhe is not None:
            _ordenar(detalhe).to_excel(writer, sheet_name="Composição das rubricas", index=False)
        if classificacao is not None:
            classificacao.to_excel(writer, sheet_name="Classificação", index=False)
        if ocorrencias is not None and not ocorrencias.empty:
            _ordenar(ocorrencias).to_excel(writer, sheet_name="Validação e ocorrências", index=False)
        for ws in writer.book.worksheets:
            if maternidade and ws.title == "Comparativo":
                continue
            _ajustar_planilha(ws)
            for celula in ws[1]:
                if any(t in str(celula.value).lower() for t in ["valor", "diferença", "crédito", "declarado", "total identificado", "principal", "apoios considerados"]):
                    for item in ws.iter_cols(
                        min_col=celula.column,
                        max_col=celula.column,
                        min_row=2,
                    ):
                        for valor in item:
                            valor.number_format = 'R$ #,##0.00'
        if maternidade:
            writer.book["Levantamento"].sheet_state = "hidden"
    return memoria.getvalue()

