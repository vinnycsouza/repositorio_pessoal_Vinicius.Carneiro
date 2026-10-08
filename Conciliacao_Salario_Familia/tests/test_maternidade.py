import io
import unittest

import openpyxl
import pandas as pd

from modules.ecac_parser import analisar_texto_dctfweb, processar_arquivos
from modules.excel_output import gerar_excel_comparativo, gerar_excel_ecac
from modules.maternidade import (
    aplicar_classificacao, comparar_maternidade, competencia, ler_manad,
    normalizar_esocial, normalizar_manad, sugerir_classificacao, validar,
    resumir_maternidade, separar_periodos_sem_dados,
)
from test_ecac_parser import TEXTO


class MaternidadeTest(unittest.TestCase):
    def base_manad(self, codigo="8006", valor="809,32", periodo="012023"):
        return pd.DataFrame([{"REG": "K300", "CNPJ/CEI": "20364206000108", "IND_FL": "1", "COD_LTC": "77", "COD_REG_TRAB": "483", "DT_COMP": periodo, "COD_RUBR": codigo, "VLR_RUBR": valor, "IND_RUBR": "P", "IND_BASE_IRRF": "1", "IND_BASE_PS": "1"}])

    def detalhe(self):
        base = pd.concat([self.base_manad(), self.base_manad("8110", "185,84")], ignore_index=True)
        normal = normalizar_manad(base, {"8006": "Situação Maternidade Empresa", "8110": "Maternidade Adicionais"}, "levantamento.xlsx")
        return validar(normal, normal, "MANAD")[0]

    def ecac(self, valor=995.16):
        return pd.DataFrame([{"cnpj": "20.364.206/0001-08", "competencia": "01/2023", "salario_maternidade": valor}])

    def test_extracao_maternidade_mensal_anual_e_ausente(self):
        texto = TEXTO.replace("Salário Família: 1.512,02", "Salário Maternidade: 995,16")
        self.assertEqual(analisar_texto_dctfweb(texto, "a.pdf").salario_maternidade, 995.16)
        anual = analisar_texto_dctfweb(texto.replace("03/2026", "2026"), "b.pdf")
        self.assertEqual(anual.competencia, "2026")
        self.assertIsNone(analisar_texto_dctfweb(TEXTO, "c.pdf").salario_maternidade)
        zero = analisar_texto_dctfweb(texto.replace("995,16", "0,00"), "d.pdf")
        self.assertEqual(zero.salario_maternidade, 0)

    def test_diferenca_unica_inclui_apoios_e_preserva_sinal(self):
        d = self.detalhe()
        d = aplicar_classificacao(d, sugerir_classificacao(d))
        r = comparar_maternidade(self.ecac(), d).iloc[0]
        self.assertEqual(r["Principal"], 809.32)
        self.assertEqual(r["Total identificado"], 995.16)
        self.assertEqual(r["Diferença"], 0)
        self.assertEqual(r["Pendência de revisão"], "")
        r = comparar_maternidade(self.ecac(700), d).iloc[0]
        self.assertEqual(r["Diferença"], 295.16)
        self.assertEqual(comparar_maternidade(self.ecac(1100), d).iloc[0]["Diferença"], -104.84)

    def test_ausencia_nao_vira_credito_zero(self):
        d = self.detalhe()
        d = aplicar_classificacao(d, sugerir_classificacao(d))
        r = comparar_maternidade(self.ecac(None), d).iloc[0]
        self.assertTrue(pd.isna(r["Diferença"]))
        ecac = pd.concat([self.ecac(), pd.DataFrame([{"cnpj": "20364206000108", "competencia": "02/2023", "salario_maternidade": 50}])])
        r = comparar_maternidade(ecac, d).iloc[1]
        self.assertTrue(pd.isna(r["Principal"]))
        self.assertTrue(pd.isna(r["Diferença"]))

    def test_repeticao_nao_e_credito_confirmado_nem_deduplicada(self):
        detalhe = self.detalhe().iloc[:1]
        repetido = pd.concat([detalhe, detalhe], ignore_index=True)
        detalhe, _ = validar(repetido, repetido, "MANAD")
        d = aplicar_classificacao(detalhe, sugerir_classificacao(detalhe))
        resultado = comparar_maternidade(self.ecac(809.32), d)
        self.assertEqual(resultado.iloc[0]["Principal"], 1618.64)
        self.assertIn("Repetição", resultado.iloc[0]["Pendência de revisão"])
        resumo = resumir_maternidade(resultado)
        self.assertEqual(resumo["Potencial crédito calculado"], 809.32)
        self.assertEqual(resumo["Desse valor, pendente por repetição"], 809.32)

    def test_sem_dados_separado_e_valor_invalido_permanece_no_comparativo(self):
        d = aplicar_classificacao(self.detalhe(), sugerir_classificacao(self.detalhe()))
        ecac = pd.concat([self.ecac(None), pd.DataFrame([
            {"cnpj": "20364206000108", "competencia": "02/2023", "salario_maternidade": None},
        ])], ignore_index=True)
        d.loc[d.grupo.eq("Principal"), "valor"] = float("nan")
        resultado = comparar_maternidade(ecac, d)
        exibicao, sem_dados = separar_periodos_sem_dados(resultado)
        self.assertEqual(list(exibicao.competencia), ["01/2023"])
        self.assertIn("inválido", exibicao.iloc[0]["Pendência de revisão"])
        self.assertEqual(list(sem_dados.competencia), ["02/2023"])
        resumo = resumir_maternidade(resultado)
        self.assertTrue(pd.isna(resumo["Potencial crédito calculado"]))
        self.assertEqual(resumo["Apurações sem dados para comparação"], 2)

    def test_credito_nao_compensa_diferencas_negativas_e_zero_e_informado(self):
        d = aplicar_classificacao(self.detalhe(), sugerir_classificacao(self.detalhe()))
        outro = d.copy()
        outro["competencia"] = "02/2023"
        ecac = pd.concat([self.ecac(0), pd.DataFrame([
            {"cnpj": "20364206000108", "competencia": "02/2023", "salario_maternidade": 1100},
        ])], ignore_index=True)
        resultado = comparar_maternidade(ecac, pd.concat([d, outro]))
        resumo = resumir_maternidade(resultado)
        self.assertEqual(resumo["Potencial crédito calculado"], 995.16)
        self.assertEqual(resumo["Apurações com diferença"], 2)
        self.assertEqual(resumo["Apurações sem dados para comparação"], 0)

    def test_validador_nao_soma_apoios_ausentes(self):
        detalhe = self.detalhe().drop(columns=["validacao", "descricao_validador", "incidencia_validador", "candidata"])
        validado, faltas = validar(detalhe.iloc[:1], detalhe, "MANAD")
        self.assertEqual(len(validado), 1)
        self.assertEqual(len(faltas), 1)
        self.assertAlmostEqual(validado.valor.sum(), 809.32)

    def test_repeticoes_preservadas_e_excesso_detectado(self):
        detalhe = self.detalhe().iloc[:1]
        excesso = pd.concat([detalhe, detalhe], ignore_index=True)
        validado, _ = validar(excesso, detalhe, "MANAD")
        self.assertEqual(len(validado), 2)
        self.assertIn("Repetição excedente", validado.iloc[1].validacao)

    def test_descricao_alternativa_e_revisao_obrigatoria(self):
        normal = normalizar_manad(self.base_manad(), {"8006": "Afastamento especial"}, "a.xlsx")
        detalhe, _ = validar(normal, normal, "MANAD")
        self.assertEqual(len(detalhe), 1)
        classes = sugerir_classificacao(detalhe)
        self.assertEqual(classes.iloc[0].grupo, "Revisar")
        with self.assertRaises(ValueError):
            aplicar_classificacao(detalhe, classes)

    def test_periodo_13_e_reclassificacao_controlada(self):
        self.assertEqual(competencia("132023"), "2023")
        base = normalizar_manad(self.base_manad(periodo="022023"), {"8006": "13º Maternidade Rescisão"}, "a.xlsx")
        detalhe, _ = validar(base, base, "MANAD")
        classes = sugerir_classificacao(detalhe)
        classes["grupo"] = "Apoio"
        with self.assertRaises(ValueError):
            aplicar_classificacao(detalhe, classes)
        classes["justificativa"] = "Revisão do período de apuração"
        classes["competencia_destino"] = "2023"
        resultado = aplicar_classificacao(detalhe, classes)
        self.assertEqual(resultado.iloc[0].competencia, "2023")
        self.assertEqual(resultado.iloc[0].competencia_original, "02/2023")

    def test_manad_cp1252_cadastro_e_multiplos_blocos(self):
        txt = "K150|20364206000108|01012023|8006|Situação Maternidade Empresa\nK300|20364206000108|1|77|483|012023|8006|809,32|P|1|1"
        fonte, _ = ler_manad(txt.encode("cp1252"), "manad.txt")
        self.assertEqual(fonte.iloc[0].descricao, "Situação Maternidade Empresa")
        self.assertEqual(fonte.iloc[0].linha_origem, 2)
        self.assertEqual(fonte.iloc[0].valor, 809.32)

    def test_esocial_identificacao_e_cadastro(self):
        base = pd.DataFrame([{"nr_insc_estab": "02633573000188", "tp_insc_estab": "1", "per_apur": "2023-01", "cod_rubr": "0014", "ide_tab_rubr": "RH", "dsc_rubr": "Salário maternidade", "vr_rubr": 100, "arquivo": "evento.xml", "caminho_item": "item[1]", "status_auditoria": "S1010_VALIDO", "cod_inc_cp": "21"}])
        fonte = normalizar_esocial(base, "incidencia.xlsx")
        detalhe, faltas = validar(fonte, fonte, "Relatórios eSocial")
        self.assertEqual(detalhe.iloc[0].validacao, "Lançamento confere")
        self.assertEqual(detalhe.iloc[0].incidencia_validador, "21")
        self.assertTrue(faltas.empty)
        fonte.loc[0, "cnpj"] = "20364206000108"
        with self.assertRaises(ValueError):
            validar(detalhe, fonte, "Relatórios eSocial")

    def test_exportacao_maternidade_e_composicao(self):
        d = self.detalhe()
        classes = sugerir_classificacao(d)
        d = aplicar_classificacao(d, classes)
        resultado = comparar_maternidade(self.ecac(), d)
        dados = gerar_excel_comparativo(resultado, self.ecac(), d, d, classes, pd.DataFrame())
        w = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
        self.assertIn("Composição das rubricas", w.sheetnames)
        self.assertIn("Classificação", w.sheetnames)
        self.assertEqual(w["Comparativo"]["B13"].value, 809.32)
        self.assertEqual([c.value for c in w["Comparativo"][12]], ["Competência", "Principal", "Apoios", "Total identificado", "Declarado no e-CAC", "Diferença", "Pendência de revisão"])
        self.assertNotIn("Critérios", w.sheetnames)
        self.assertEqual(w["Levantamento"].sheet_state, "hidden")
        self.assertEqual(w["Comparativo"]["C4"].value, 0)
        ecac = gerar_excel_ecac(self.ecac(), pd.DataFrame(), pd.DataFrame())
        self.assertIn("Salário Maternidade", openpyxl.load_workbook(io.BytesIO(ecac)).sheetnames)

    def test_todas_abas_exportadas_em_ordem_cronologica(self):
        base = pd.DataFrame([{"cnpj": "20364206000108", "competencia": c, "valor": i} for i, c in enumerate(["01/2023", "2022", "12/2021", "01/2022", "12/2022", "2021"])])
        esperado = ["12/2021", "2021", "01/2022", "12/2022", "2022", "01/2023"]
        dados = gerar_excel_comparativo(base, base, base, base, ocorrencias=base)
        w = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
        for nome in ["Comparativo", "Dados e-CAC", "Levantamento", "Composição das rubricas", "Validação e ocorrências"]:
            with self.subTest(aba=nome):
                self.assertEqual([r[1] for r in list(w[nome].values)[1:]], esperado)
        self.assertEqual(base.competencia.iloc[0], "01/2023")

    def test_exportacao_multiplos_cnpjs_preserva_identificacao(self):
        d = aplicar_classificacao(self.detalhe(), sugerir_classificacao(self.detalhe()))
        outro = d.copy()
        outro["cnpj"] = "02633573000188"
        ecac = pd.concat([self.ecac(), pd.DataFrame([
            {"cnpj": "02633573000188", "competencia": "01/2023", "salario_maternidade": 995.16},
        ])], ignore_index=True)
        composto = pd.concat([d, outro], ignore_index=True)
        resultado = comparar_maternidade(ecac, composto)
        arquivo = gerar_excel_comparativo(resultado, ecac, composto, composto)
        ws = openpyxl.load_workbook(io.BytesIO(arquivo), data_only=True)["Comparativo"]
        self.assertEqual(ws["A12"].value, "CNPJ")
        self.assertEqual({ws["A13"].value, ws["A14"].value}, {"02633573000188", "20364206000108"})


if __name__ == "__main__":
    unittest.main()
