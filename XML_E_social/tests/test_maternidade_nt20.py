import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from modules.auditoria import (
    classificar_status_cp, entra_base_cp, preparar_movimentos_cp,
    revisar_maternidade_dataframe, _filtrar_movimentos_cp_exportacao,
)
from modules.data_source import WorkspaceContext, SQLiteDataSource
from modules.levantamento_sqlite import FiltrosLevantamento, consultar_rubricas
from modules.processador_zip import processar_fontes_esocial, preparar_workspace_para_relatorios
from modules.sqlite_relatorio import atualizar_regra_maternidade, gerar_excel_saida_sqlite
from test_carga_incremental import S1010, S1200, zip_xmls


class MaternidadeNT20Test(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        env = patch.dict(os.environ, {"ESOCIAL_WORKSPACES_DIR": self.temp.name})
        env.start()
        self.addCleanup(env.stop)

    def _workspace(self):
        original = S1200[S1200.index("<itensRemun>"):S1200.index("</itensRemun>") + len("</itensRemun>")]
        itens = "".join(original.replace("<codRubr>100</codRubr>", f"<codRubr>{codigo}</codRubr>")
                        .replace("100.00", str(valor))
                        for codigo, valor in [("100", 100), ("200", 30), ("300", 40), ("400", 10), ("500", 5)])
        arquivos = {"s1200.xml": S1200.replace(original, itens)}
        for codigo, cp in [("100", "11"), ("200", "21"), ("300", "22"), ("400", "00"), ("500", "99")]:
            arquivos[codigo + ".xml"] = (S1010.replace('Id="ID1010"', f'Id="ID{codigo}"')
                .replace("<codRubr>100</codRubr>", f"<codRubr>{codigo}</codRubr>")
                .replace("<codIncCP>11</codIncCP>", f"<codIncCP>{cp}</codIncCP>"))
        return processar_fontes_esocial([("sintetico.zip", zip_xmls(**arquivos))])

    def test_classificacao_nao_presume_patronal_ou_credito(self):
        for codigo in ("21", "22"):
            self.assertEqual(classificar_status_cp(codigo), "Revisar codIncCP")
            self.assertFalse(entra_base_cp(codigo))
        for codigo in ("11", "12", "13", "14", "15", "16", "23", "24", "25", "26"):
            self.assertEqual(classificar_status_cp(codigo), "Incide CP")
        self.assertEqual(classificar_status_cp("00"), "Não incide CP")
        self.assertEqual(classificar_status_cp(""), "Sem S-1010")
        self.assertEqual(classificar_status_cp("99"), "Revisar codIncCP")
        df = pd.DataFrame([{"cod_inc_cp": cp, "vr_rubr": 30, "per_apur": periodo}
                           for cp in ("21", "22")
                           for periodo in ("2019-01", "2020-11", "2020-12", "2026", "2026-01")])
        revisado = preparar_movimentos_cp(df)
        self.assertTrue(revisado["status_cp"].eq("Revisar codIncCP").all())
        self.assertTrue(revisado["considerado_cp"].eq("Não").all())
        self.assertTrue(revisado["observacao"].str.contains("NT 20/2020", regex=False).all())
        self.assertEqual(sorted(revisado["per_apur"]), sorted(df["per_apur"]))
        self.assertEqual(revisado["vr_rubr"].sum(), df["vr_rubr"].sum())

    def test_excel_legado_preserva_observacao_e_filtro(self):
        df = pd.DataFrame({"cod_inc_cp": ["11", "21", "22", "00"],
                           "status_cp": ["Incide CP"] * 3 + ["Não incide CP"],
                           "considerado_cp": ["Sim"] * 3 + ["Não"],
                           "observacao": ["origem"] * 4, "vr_rubr": [100, 30, 40, 10]})
        novo = revisar_maternidade_dataframe(df)
        self.assertEqual(df.loc[1, "status_cp"], "Incide CP")
        self.assertTrue(novo.loc[1, "observacao"].startswith("origem | "))
        pd.testing.assert_frame_equal(novo, revisar_maternidade_dataframe(novo))
        self.assertEqual(_filtrar_movimentos_cp_exportacao(df, "incidencia_cp_padrao")["cod_inc_cp"].tolist(), ["11"])
        self.assertEqual(_filtrar_movimentos_cp_exportacao(df, "analise_prioritaria")["cod_inc_cp"].tolist(), ["11", "21", "22"])
        sem_status = df.drop(columns=["status_cp", "considerado_cp"])
        self.assertEqual(_filtrar_movimentos_cp_exportacao(sem_status, "incidencia_cp_padrao")["cod_inc_cp"].tolist(), ["11"])

    def test_workspace_novo_relatorio_e_levantamento(self):
        resultado = self._workspace()
        c = sqlite3.connect(resultado["db_path"])
        try:
            self.assertEqual(c.execute("SELECT total_incide_cp,total_s1200 FROM rel_base_trabalhador").fetchone(), (100, 185))
            self.assertEqual(c.execute("SELECT COUNT(*) FROM rel_movimentos_cp").fetchone()[0], 5)
        finally:
            c.close()
        fonte = SQLiteDataSource(WorkspaceContext.from_path(resultado["db_path"]))
        incide = consultar_rubricas(fonte, FiltrosLevantamento(status_cp="Incide CP"))
        self.assertEqual(set(incide["cod_rubr"]), {"100"})
        revisao = consultar_rubricas(fonte, FiltrosLevantamento(status_cp="Revisar codIncCP"))
        self.assertEqual(set(revisao["cod_rubr"]), {"200", "300", "500"})
        for modo, codigos in [("incidencia_cp_padrao", {"100"}),
                              ("analise_prioritaria", {"100", "200", "300", "500"})]:
            destino = Path(self.temp.name) / (modo + ".xlsx")
            gerar_excel_saida_sqlite(resultado["db_path"], destino, resultado["empresa"],
                                     resultado["pacote_sqlite"]["resumo_visual"],
                                     resultado["pacote_sqlite"]["rubricas_cp"],
                                     modo_exportacao_movimentos_cp=modo)
            wb = load_workbook(destino, read_only=True, data_only=True)
            try:
                for aba in ("03_movimentos_cp", "apoio_s1200"):
                    linhas = wb[aba].iter_rows(values_only=True)
                    cabecalho = next(linhas)
                    dados = [dict(zip(cabecalho, linha)) for linha in linhas]
                    self.assertEqual({r["cod_rubr"] for r in dados}, codigos)
                    for r in dados:
                        if r["cod_inc_cp"] in ("21", "22"):
                            self.assertIn("NT 20/2020", r["observacao_cp_patronal"])
                            self.assertEqual(r["considerado_cp"], "Não")
            finally:
                wb.close()

    def test_migracao_legada_atomica_idempotente_sem_reler_xml(self):
        resultado = self._workspace()
        c = sqlite3.connect(resultado["db_path"])
        try:
            originais = c.execute("SELECT * FROM eventos").fetchall()
            objetos = c.execute("SELECT * FROM objetos").fetchall()
            for tabela in ("dados_remuneracoes", "rel_movimentos_cp", "rel_rubricas_cp_base"):
                c.execute(f"UPDATE {tabela} SET status_cp='Incide CP',considerado_cp='Sim' WHERE cod_inc_cp IN ('21','22')")
            # S-2299 segue separado; só a parcela de maternidade é retirada
            # da reconciliação em que anteriormente havia sido considerada.
            for tabela in ("dados_remuneracoes_s2299", "rel_movimentos_s2299"):
                c.execute(
                    f"INSERT INTO {tabela}(per_apur,nr_insc_estab,cod_lotacao,cod_categ,"
                    "cod_inc_cp,status_cp,considerado_cp,vr_rubr) "
                    "VALUES('2026-01','12345678000199','1','101','22','Incide CP','Sim',20)"
                )
            c.execute("UPDATE rel_base_trabalhador SET total_incide_cp=170,diferenca_incide_cp_vs_s5001=170")
            c.execute("UPDATE rel_reconciliacao_s5011 SET base_s1200=170,base_s2299=20,base_s5011=100,base_teorica_total=190,diferenca_s1200_s5011=70,diferenca_total_s5011=90,status_reconciliacao='Revisar'")
            c.execute("DELETE FROM meta WHERE chave='versao_regra_maternidade'")
            c.commit()
            # Uma falha não pode deixar bases novas e movimentos antigos.
            c.execute("CREATE TEMP TRIGGER falha_nt20 BEFORE UPDATE ON rel_movimentos_cp BEGIN SELECT RAISE(ABORT,'falha simulada'); END")
            with self.assertRaises(sqlite3.IntegrityError):
                atualizar_regra_maternidade(c)
            self.assertEqual(c.execute("SELECT total_incide_cp FROM rel_base_trabalhador").fetchone()[0], 170)
            self.assertIsNone(c.execute("SELECT valor FROM meta WHERE chave='versao_regra_maternidade'").fetchone())
            c.execute("DROP TRIGGER falha_nt20")
            c.commit()
            with patch('modules.processador_zip.ET.fromstring', side_effect=AssertionError('XML não pode ser relido')):
                preparar_workspace_para_relatorios(resultado["db_path"])
            self.assertEqual(c.execute("SELECT total_incide_cp,diferenca_incide_cp_vs_s5001 FROM rel_base_trabalhador").fetchone(), (100, 100))
            self.assertEqual(c.execute("SELECT base_s1200,base_teorica_total,diferenca_total_s5011,status_reconciliacao FROM rel_reconciliacao_s5011").fetchone(), (100, 100, 0, 'OK'))
            self.assertEqual(c.execute("SELECT * FROM eventos").fetchall(), originais)
            self.assertEqual(c.execute("SELECT * FROM objetos").fetchall(), objetos)
            self.assertEqual(c.execute("SELECT base_s2299 FROM rel_reconciliacao_s5011").fetchone()[0], 0)
            self.assertEqual(c.execute("SELECT vr_rubr,status_cp,considerado_cp FROM rel_movimentos_s2299").fetchone(), (20, 'Revisar codIncCP', 'Não'))
            self.assertEqual(c.execute("SELECT COUNT(*) FROM rel_movimentos_cp").fetchone()[0], 5)
            self.assertEqual(c.execute("SELECT soma_valores FROM rel_controle_integridade WHERE tabela='rel_reconciliacao_s5011'").fetchone()[0], 0)
            self.assertIsNotNone(c.execute("SELECT valor FROM meta WHERE chave='backup_classificacao_maternidade'").fetchone())
            self.assertFalse(atualizar_regra_maternidade(c))
            self.assertEqual(c.execute("SELECT total_incide_cp FROM rel_base_trabalhador").fetchone()[0], 100)
        finally:
            c.close()

    def test_migracao_nao_altera_carga_interrompida(self):
        resultado = self._workspace()
        c = sqlite3.connect(resultado["db_path"])
        try:
            c.execute("DELETE FROM meta WHERE chave='versao_regra_maternidade'")
            c.execute("UPDATE meta SET valor='interrompido' WHERE chave='status'")
            c.commit()
            antes = c.execute('SELECT * FROM rel_movimentos_cp').fetchall()
            with self.assertRaisesRegex(ValueError, 'Conclua ou retome'):
                preparar_workspace_para_relatorios(resultado["db_path"])
            self.assertEqual(c.execute('SELECT * FROM rel_movimentos_cp').fetchall(), antes)
            self.assertIsNone(c.execute("SELECT valor FROM meta WHERE chave='versao_regra_maternidade'").fetchone())
        finally:
            c.close()
