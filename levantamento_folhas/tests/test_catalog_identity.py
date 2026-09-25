import io,unittest
from unittest.mock import patch,Mock
from types import SimpleNamespace
import openpyxl
import test_core as fixtures
import core

class CatalogIdentityTests(unittest.TestCase):
    def load(self,identity=None,company_column=None,manual='',sheet='apoio_s1010'):
        headers=['cod_rubr','ide_tab_rubr','dsc_rubr','cod_inc_cp','tp_rubr','ini_valid']
        rows=[['0001','RH3','SALARIO','11','1','2018-07']]
        if company_column is not None:
            headers+=['cnpj_empregador']
            rows=[rows[0]+[x] for x in company_column]
        class Book(dict):
            close=Mock()
        book=Book({sheet:SimpleNamespace(iter_rows=lambda **kw:iter([headers]+rows))})
        if identity is not None:book['00_empresa']=SimpleNamespace(iter_rows=lambda **kw:iter([['cnpj_empregador']]+[[x] for x in identity]))
        with patch.object(core,'load_catalog_workbook',return_value=book):
            return core.import_catalog('test',manual)
    def test_existing_identification_preserved(self):
        r=self.load(identity=['02.633.573/0001-88'])
        self.assertEqual(r['empresa_raiz'],'02633573');self.assertFalse(r['identificacao_manual'])
    def test_missing_sheet_uses_explicit_rubric_company(self):
        r=self.load(company_column=['02633573000188'])
        self.assertEqual(r['empresa_raiz'],'02633573');self.assertIn('apoio_s1010',r['identificacao_empresa'])
    def test_missing_identity_requires_input(self):
        with self.assertRaisesRegex(ValueError,'Preencha o campo CNPJ'):self.load()
    def test_manual_binding_and_fallback_sheet(self):
        r=self.load(manual='02.633.573/0001-88',sheet='02_rubricas_cp')
        self.assertTrue(r['identificacao_manual']);self.assertEqual(r['rubricas'][0]['cod_rubr'],'0001')
        self.assertEqual(r['fonte'],'02_rubricas_cp')
    def test_empty_identity_sheet_accepts_manual(self):
        self.assertTrue(self.load(identity=[],manual='02633573')['identificacao_manual'])
    def test_conflicts_cannot_be_overridden(self):
        with self.assertRaisesRegex(ValueError,'diverge'):self.load(identity=['02633573'],manual='12345678')
        with self.assertRaisesRegex(ValueError,'empresas diferentes'):self.load(company_column=['02633573','12345678'],manual='02633573')
        with self.assertRaisesRegex(ValueError,'empresas diferentes'):self.load(identity=['02633573'],company_column=['12345678'])
    def test_branches_same_root_are_allowed(self):
        self.assertEqual(self.load(company_column=['02633573000188','02633573000200'])['empresa_raiz'],'02633573')
    def test_manual_identity_keeps_company_gate_and_export_disclosure(self):
        f=fixtures.CoreTests();f.setUp();f.a['catalog']=self.load(manual='12345678')
        self.assertEqual(core.details(f.a)[0]['correspondencia'],'Empresa diferente')
        w=openpyxl.load_workbook(io.BytesIO(core.export_excel(f.a)))
        self.assertTrue(any('vinculada manualmente' in str(r[0]) for r in w.worksheets[0].values))
    def test_bad_manual_cnpj_is_not_silently_padded(self):
        with self.assertRaisesRegex(ValueError,'inválida'):self.load(manual='12345')

if __name__=='__main__':unittest.main()

