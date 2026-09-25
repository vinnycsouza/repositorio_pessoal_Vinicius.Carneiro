import copy
import io
import unittest
from decimal import Decimal,ROUND_HALF_UP
import test_twenty as fixtures
import core
import openpyxl

class ProportionalTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.TwentyTests();self.f.setUp();self.f.mixed()
        self.a=self.f.a;self.d=self.f.d

    def test_estimate_is_independent_of_existing_hypothesis(self):
        original=copy.deepcopy(self.a);before=core.twenty_composition(self.a)
        s,e=core.proportional_twenty(self.a)
        expected=int((Decimal(3000000)*Decimal(3000000)/Decimal(3100000)).quantize(Decimal(1),rounding=ROUND_HALF_UP))
        self.assertEqual(s[0]['saldo_estimado_centavos'],expected)
        self.assertEqual(s[0]['diferenca_para_base_20_centavos'],3000000-expected)
        self.assertEqual(core.twenty_composition(self.a),before)
        self.assertEqual(self.a,original)

    def test_reductions_and_conflicting_candidates_stay_separate(self):
        self.f.f.add('0870','ATRASO',10000,'11','Desconto')
        r,c=self.f.f.add('0999','BONUS',20000,'11')
        self.f.f.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        s,e=core.proportional_twenty(self.a)
        self.assertEqual(s[0]['quantidade_parcelas'],2)
        self.assertEqual(s[0]['quantidade_candidatas'],1)
        self.assertGreater(s[0]['reducoes_estimadas_centavos'],0)
        self.assertEqual(s[0]['saldo_estimado_centavos'],s[0]['acrescimos_estimados_centavos']-s[0]['reducoes_estimadas_centavos'])
        self.assertEqual(sum(r['estimativa_centavos'] for r in e if r['cenario']=='Estimativa pelo cadastro'),s[0]['saldo_estimado_centavos'])

    def test_invalid_zero_or_unmixed_references_do_not_estimate(self):
        self.d['grupos_empresa'][0]['quantidade_base']=3
        self.assertEqual(core.proportional_twenty(self.a),([],[]))
        self.d['grupos_empresa'][0]['quantidade_base']=2
        self.d['grupos_empresa']=self.d['grupos_empresa'][:1];self.d['bases']['mensal']=3000000
        self.assertEqual(core.proportional_twenty(self.a),([],[]))
        self.d['grupos_empresa'][0].update(grupo='Contribuição zerada',contribuicao_centavos=0)
        self.assertEqual(core.proportional_twenty(self.a),([],[]))

    def test_monthly_and_thirteenth_have_different_factors(self):
        self.f.f.add('0290','13 SALARIO',10000,'12')
        self.d['bases']['13']=20000
        self.d['grupos_empresa'] += [
            {'base':'13º','grupo':'Alíquota de 20%','quantidade_base':1,'aliquota_percentual':20,'base_centavos':10000,
             'contribuicao_centavos':2000,'linha_contribuicao':'Valor prev. empresa sobre 13º Salário 1 20% 20,00','pagina':1},
            {'base':'13º','grupo':'Contribuição zerada','quantidade_base':1,'aliquota_percentual':None,'base_centavos':10000,
             'contribuicao_centavos':0,'linha_contribuicao':'Valor prev. empresa sobre 13º Salário 0 0,00','pagina':1}]
        s,e=core.proportional_twenty(self.a)
        self.assertEqual(s[1]['fator_proporcao'],0.5)
        self.assertEqual(s[1]['saldo_estimado_centavos'],5000)
        self.assertNotEqual(s[0]['fator_proporcao'],s[1]['fator_proporcao'])

    def test_excel_numeric_percent_and_money(self):
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        s=w['Simulacao proporcional'];heads=[c.value for c in s[1]]
        c=s.cell(2,heads.index('Proporção monetária dos 20%')+1)
        self.assertEqual(c.data_type,'n');self.assertEqual(c.number_format,'0.00%')
        self.assertEqual(s.cell(2,heads.index('Saldo estimado (R$)')+1).data_type,'n')
