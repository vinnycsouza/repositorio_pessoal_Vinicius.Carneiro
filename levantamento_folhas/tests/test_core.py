import io, sys, unittest, zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import core
import openpyxl

class CoreTests(unittest.TestCase):
    def setUp(self):
        self.d={'cnpj':'02.633.573/0001-88','competencia':'2024-03','tipo':'Mensal','hash':'x','arquivo':'teste.pdf','empresa':'AJ','bases':{'mensal':90000},'totais':{},'checagens':[],'paginas':1,'alertas':[], 'rubricas':[
            {'codigo':'0001','descricao':'SALARIO','lado':'Provento','valor_centavos':100000,'pagina':1},
            {'codigo':'0116','descricao':'FALTAS','lado':'Desconto','valor_centavos':10000,'pagina':1}]}
        self.cat={'empresa_raiz':'02633573','rubricas':[{'cod_rubr':'0001','ide_tab_rubr':'RH3','dsc_rubr':'SALARIO','cod_inc_cp':'11','tp_rubr':'1','ini_valid':'2018-07','fim_valid':''},{'cod_rubr':'0116','ide_tab_rubr':'RH3','dsc_rubr':'FALTAS','cod_inc_cp':'11','tp_rubr':'2','ini_valid':'2018-07','fim_valid':''}]}
        self.a={'name':'Teste','docs':[self.d],'catalog':self.cat,'decisions':{},'errors':[]}
    def test_discount_reduces_and_suggestion_is_not_confirmation(self):
        c=core.reconciliation(self.a)[0]
        self.assertEqual(c['reconstruida_centavos'],90000)
        self.assertEqual(c['diferenca_centavos'],0)
        self.assertEqual(c['estado'],'Hipótese com sugestões')
    def test_old_period_is_projected(self):
        self.d['competencia']='2012-08'
        self.assertEqual(core.details(self.a)[0]['efeito'],'Acrescenta (sugestão)')
    def test_other_company_is_pending(self):
        self.d['cnpj']='08.362.490/0001-88'
        self.assertEqual(core.details(self.a)[0]['efeito'],'Pendente')
    def test_maternity_not_ordinary_addition(self):
        self.cat['rubricas'][0]['cod_inc_cp']='21'
        self.assertEqual(core.details(self.a)[0]['efeito'],'Pendente')
    def test_ambiguous_history_is_pending(self):
        self.cat['rubricas'].append({**self.cat['rubricas'][0],'cod_inc_cp':'00'})
        self.assertEqual(core.details(self.a)[0]['efeito'],'Pendente')
    def test_excel_numeric_and_identifiers(self):
        b=core.export_audit_excel(self.a); w=openpyxl.load_workbook(io.BytesIO(b),data_only=False)
        s=w['Cruzamento por folha']; heads=[c.value for c in s[1]]
        self.assertEqual(s.cell(2,heads.index('Código')+1).value,'0001')
        self.assertEqual(s.cell(2,heads.index('Valor integral da rubrica (R$)')+1).value,1000)
        self.assertIn('Conferencia bases',w.sheetnames)
    def test_zip_paths_are_not_written(self):
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as z: z.writestr('../../outside.pdf',b'%PDF-test')
        items=list(core.unpack('test.zip',b.getvalue()))
        self.assertEqual(len(items),1)
        self.assertEqual(items[0][1],b'%PDF-test')
    def test_previdencia_missing_and_zero_are_distinct(self):
        self.d['previdencia']={'previdencia_empresa_total':{'valor_centavos':0,'pagina':2}}
        rows=core.previdencia_rows(self.d)
        self.assertEqual(len(rows),4)
        self.assertIsNone(rows[0]['valor_centavos'])
        self.assertEqual(rows[1]['valor_centavos'],0)
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        s=w['Resumo previdenciario']
        self.assertIsNone(s['E2'].value)
        self.assertEqual(s['E3'].value,0)

    def test_no_automatic_interest_selection(self):
        rows=core.details(self.a)
        self.assertTrue(all(not r['selecionada'] for r in rows))
        self.assertEqual(rows[1]['grupo'],'Reduções da base')

    def test_historical_reference_projects_with_explicit_label(self):
        self.d['competencia']='2012-08'
        row=core.details(self.a)[0]
        self.assertEqual(row['codIncCP'],'11')
        self.assertEqual(row['descricao_relatorio'],'SALARIO')
        self.assertEqual(row['referencia'],'Projeção pelo cadastro disponível')
        self.assertEqual(row['grupo'],'Possíveis acréscimos')

    def test_selection_does_not_change_effect(self):
        row=core.details(self.a)[1]
        self.a['selections']={core.selection_key(row):True}
        after=core.details(self.a)[1]
        self.assertTrue(after['selecionada'])
        self.assertEqual(after['grupo'],'Reduções da base')
        self.assertEqual(core.reconciliation(self.a)[0]['reconstruida_centavos'],90000)

    def test_brazilian_display_and_numeric_excel(self):
        self.assertEqual(core.brl(253071630),'R$ 2.530.716,30')
        self.assertEqual(core.brl(-30000),'R$ -300,00')
        self.assertEqual(core.brl(0),'R$ 0,00')
        self.assertEqual(core.brl(None),'Não informado')
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        s=w['Cruzamento por folha']; headings=[c.value for c in s[1]]
        cell=s.cell(2,headings.index('Valor integral da rubrica (R$)')+1)
        self.assertEqual(cell.data_type,'n')
        self.assertIn('R$',cell.number_format)
        self.assertEqual(s.max_row,3)
        groups=[s.cell(i,headings.index('Grupo')+1).value for i in (2,3)]
        self.assertEqual(groups,['Possíveis acréscimos','Reduções da base'])

    def test_previdencia_filter_distinguishes_missing_zero_positive(self):
        docs=[{'previdencia':{'previdencia_empresa_total':{'valor_centavos':v}}} for v in [None,0,123,-1]]+[{}]
        self.assertEqual(len(core.filter_previdencia(docs,'Todas')),5)
        self.assertEqual(len(core.filter_previdencia(docs,'Não localizado')),2)
        self.assertEqual(len(core.filter_previdencia(docs,'Igual a zero')),1)
        self.assertEqual(len(core.filter_previdencia(docs,'Maior que zero')),1)
        self.assertEqual(len(core.filter_previdencia(docs,'Menor que zero')),1)
        self.assertEqual(len(docs),5)

    def test_historical_conflict_prevents_projection(self):
        self.d['competencia']='2012-08'
        self.cat['rubricas'].append({**self.cat['rubricas'][0],'ini_valid':'2024-01','cod_inc_cp':'00'})
        row=core.details(self.a)[0]
        self.assertEqual(row['grupo'],'Não determinado')
        self.assertEqual(row['correspondencia'],'Conflito no histórico')

    def test_nearest_historical_reference_is_selected(self):
        self.d['competencia']='2012-08'
        self.cat['rubricas'].append({**self.cat['rubricas'][0],'ini_valid':'2024-01'})
        row=core.details(self.a)[0]
        self.assertTrue(row['vigencia_relatorio'].startswith('2018-07'))
        self.assertEqual(row['efeito'],'Acrescenta (sugestão)')

    def test_description_mismatch_does_not_project(self):
        self.d['competencia']='2012-08';self.d['rubricas'][0]['descricao']='BONUS'
        self.assertEqual(core.details(self.a)[0]['grupo'],'Não determinado')

    def test_missing_base_is_not_zero(self):
        self.d['bases']={}
        c=core.reconciliation(self.a)[0]
        self.assertIsNone(c['informada_centavos']); self.assertIsNone(c['diferenca_centavos'])


class GroupTests(unittest.TestCase):
    def parse(self,right):
        from unittest.mock import MagicMock,patch
        words=[]
        for offset,text in [(10,'Base empresa 32 6.557,74'),(310,right)]:
            for i,part in enumerate(text.split()): words.append({'text':part,'x0':offset+i*20,'top':100})
        page=MagicMock();page.width=600;page.extract_words.return_value=words
        pdf=MagicMock();pdf.pages=[page]
        with patch.object(core.pdfplumber,'open') as op:
            op.return_value.__enter__.return_value=pdf
            return core.extract_company_groups(b'')[0]
    def test_zero_is_explicit(self):
        r=self.parse('Valor da previdência empresa 0 0,00')
        self.assertEqual(r['grupo'],'Contribuição zerada')
        self.assertEqual(r['quantidade_base'],32)
        self.assertEqual(r['contribuicao_centavos'],0)
    def test_missing_is_not_zero(self):
        r=self.parse('')
        self.assertEqual(r['grupo'],'Não determinado')
        self.assertIsNone(r['contribuicao_centavos'])
    def test_twenty_and_other_rates(self):
        self.assertEqual(self.parse('Valor da previdência empresa 32 20% 1.311,55')['grupo'],'Alíquota de 20%')
        self.assertEqual(self.parse('Valor da previdência empresa 32 10% 655,77')['grupo'],'Outra alíquota')
    def test_summary_missing_is_not_zero(self):
        self.assertIsNone(core.company_summary({'bases':{}})[0]['base_20_centavos'])

if __name__=='__main__': unittest.main()
