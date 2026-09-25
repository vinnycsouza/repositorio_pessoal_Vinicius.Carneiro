import copy,io,unittest
import openpyxl
import core,simple_report
import test_twenty as fixtures

class SimpleReportTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.TwentyTests();self.f.setUp();self.f.mixed()
        self.a=self.f.a;self.d=self.f.d
    def book(self,a=None):return openpyxl.load_workbook(io.BytesIO(core.export_excel(a or self.a)))
    def test_three_sheets_and_scenarios_keep_their_own_totals(self):
        w=self.book();self.assertEqual(w.sheetnames,list(simple_report.SHEETS))
        rows=list(w.worksheets[0].values)
        amounts={r[0]:r[3] for r in rows if r[0]}
        self.assertEqual(amounts['Total projetado / hipótese'],30000)
        expected=core.proportional_twenty(self.a)[0][0]['saldo_estimado_centavos']/100
        self.assertNotIn('Total estimado líquido',amounts)
        simulation={r[0]:r[3] for r in w['Simulação proporcional'].values if r[0]}
        self.assertEqual(simulation['Total estimado líquido'],expected)
        self.assertNotEqual(expected,30000)
        self.assertIn('Estimativa proporcional — cenário independente, a confirmar na folha',simulation)
    def test_single_month_is_same_block_as_consolidated(self):
        other=copy.deepcopy(self.d);other.update(hash='earlier',competencia='2014-01')
        a={**self.a,'docs':[self.d,other]}
        full=self.book(a);month=self.book()
        for name in simple_report.SHEETS:
            start=next(i for i,r in enumerate(list(full[name].values)) if str(r[0]).startswith('2015-07 ·'))
            short=next(i for i,r in enumerate(list(month[name].values)) if str(r[0]).startswith('2015-07 ·'))
            self.assertEqual(list(full[name].values)[start:],list(month[name].values)[short:])
    def test_zero_absence_numeric_and_source_identifiers(self):
        self.d['previdencia']={'previdencia_empresa_total':{'valor_centavos':0,'pagina':1}}
        w=self.book();s=w.worksheets[1]
        amounts={r[0]:r[3] for r in s.values if r[0]}
        self.assertEqual(amounts['Total de previdência empresa'],0)
        self.assertEqual(amounts['Total da base empresa'],'Não disponível')
        rubric=next(r for r in s if r[0].value=='0012')
        self.assertEqual(rubric[0].data_type,'s');self.assertEqual(rubric[2].data_type,'n')
    def test_candidates_and_negative_reductions_remain_separate(self):
        self.f.f.add('0870','ATRASO',10000,'11','Desconto')
        r,c=self.f.f.add('0999','BONUS',20000,'11')
        self.f.f.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        w=self.book();s=w.worksheets[1]
        rubric=next(r for r in s.values if r[0]=='0870')
        self.assertEqual(rubric[3],-100)
        self.assertTrue(any(str(r[0]).startswith('Candidatas conflitantes') for r in s.values))
    def test_does_not_mutate_analysis_and_scopes_documents(self):
        before=copy.deepcopy(self.a);self.book();self.assertEqual(self.a,before)
        self.assertEqual(simple_report.select_documents(self.a['docs'],[self.d['cnpj']],['2015-07'],['Mensal']),self.a['docs'])
        self.assertEqual(simple_report.select_documents(self.a['docs'],[self.d['cnpj']],['2014-01'],['Mensal']),[])
    def test_inconsistent_reference_does_not_estimate(self):
        self.d['grupos_empresa'][0]['quantidade_base']=3
        rows=list(self.book().worksheets[0].values)
        amount=next(r for r in rows if r[0]=='Base dos 20% informada no PDF')
        self.assertEqual(amount[3],'Não disponível')
        self.assertFalse(any(str(r[0]).startswith('Estimativa proporcional') for r in rows))
    def test_source_formula_text_is_literal(self):
        self.d['rubricas'][0]['descricao']='=1+1'
        s=self.book().worksheets[1]
        c=next(r[1] for r in s if r[1].value=='=1+1')
        self.assertEqual(c.data_type,'s')
    def test_empty_scope(self):
        w=self.book({**self.a,'docs':[]})
        self.assertEqual(len(w.sheetnames),3)
        self.assertTrue(any(r[0]=='Nenhuma folha no recorte selecionado.' for r in w.worksheets[0].values))

if __name__=='__main__':unittest.main()
