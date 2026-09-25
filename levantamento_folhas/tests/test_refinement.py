import io
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import core
import openpyxl


class RefinementTests(unittest.TestCase):
    def setUp(self):
        self.d={'cnpj':'02.633.573/0001-88','empresa':'AJ','competencia':'2015-07','tipo':'Mensal',
                'hash':'test','arquivo':'teste.pdf','bases':{'13':2932801},'totais':{},'checagens':[],
                'paginas':1,'alertas':[],'rubricas':[]}
        self.cat={'empresa_raiz':'02633573','rubricas':[]}
        self.a={'name':'Teste','docs':[self.d],'catalog':self.cat,'decisions':{},'errors':[]}

    def add(self,code,name,value,cp='12',side='Provento',qty='1'):
        r={'codigo':code,'descricao':name,'lado':side,'valor_centavos':value,'quantidade':qty,'pagina':1}
        c={'cod_rubr':code,'ide_tab_rubr':'RH3','dsc_rubr':name,'cod_inc_cp':cp,
           'tp_rubr':'1' if side=='Provento' else '2','ini_valid':'2018-07','fim_valid':''}
        self.d['rubricas'].append(r);self.cat['rubricas'].append(c)
        return r,c

    def test_whitespace_does_not_change_original_key(self):
        r,c=self.add('0110','HS.EXTRAS 50%',10000,'11')
        original=core.key_for(self.d,r);c['dsc_rubr']='HS.EXTRAS  50%'
        row=core.details(self.a)[0]
        self.assertEqual(row['grupo'],'Possíveis acréscimos')
        self.assertEqual(row['chave'],original)
        self.assertEqual(row['descricao'],'HS.EXTRAS 50%')

    def test_description_variant_requires_exact_identity_and_fiscal_agreement(self):
        r,c=self.add('0260','1/3 DE FERIAS',10000,'11')
        self.cat['rubricas'].append({**c,'dsc_rubr':'13 DE FERIAS'})
        self.assertEqual(core.details(self.a)[0]['grupo'],'Possíveis acréscimos')
        self.cat['rubricas'][-1]['cod_inc_cp']='00'
        self.assertEqual(core.details(self.a)[0]['grupo'],'Não determinado')

    def test_different_tables_are_not_silently_merged(self):
        r,c=self.add('0260','1/3 DE FERIAS',10000,'11')
        self.cat['rubricas'].append({**c,'ide_tab_rubr':'OUTRA'})
        self.assertEqual(core.details(self.a)[0]['grupo'],'Não determinado')

    def test_real_conflict_is_only_a_candidate_not_adopted(self):
        self.add('0290','13 SALARIO',2914682)
        r,c=self.add('0460','13 SAL INDENIZADO',18119)
        self.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        rows=core.details(self.a)
        self.assertEqual(rows[1]['efeito'],'Pendente')
        self.assertEqual(rows[1]['base'],'Não determinada')
        report=core.reconciliation(self.a)[1]
        self.assertEqual(report['reconstruida_centavos'],2914682)
        self.assertEqual(report['candidatas_centavos'],18119)
        self.assertEqual(report['saldo_apos_todas_candidatas_centavos'],0)
        self.assertEqual(report['conclusao_composicao'],'Fechamento condicionado às candidatas')
        self.assertFalse(rows[1]['selecionada'])

    def test_candidates_do_not_search_subsets_to_close(self):
        self.add('0290','13 SALARIO',2914682)
        for code,value in [('0460',18119),('0999',12345)]:
            r,c=self.add(code,code,value);self.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        report=core.reconciliation(self.a)[1]
        self.assertEqual(report['saldo_apos_todas_candidatas_centavos'],-12345)
        self.assertNotEqual(report['conclusao_composicao'],'Fechamento condicionado às candidatas')

    def test_no_catalog_does_not_default_to_monthly(self):
        self.add('9999','EVENTO',10000)
        self.a['catalog']=None
        self.assertEqual(core.details(self.a)[0]['base'],'Não determinada')

    def test_insured_contribution_is_not_a_base_reduction(self):
        r,c=self.add('0510','INSS',10000,'31','Desconto')
        self.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        row=core.details(self.a)[0]
        self.assertEqual(row['natureza'],'Contribuição do segurado')
        self.assertEqual(row['base'],'Não se aplica')
        self.assertEqual(core.composition_trace(self.a)[0]['parcela_centavos'],0)

    def test_insured_label_does_not_hide_real_base_conflict(self):
        r,c=self.add('0510','INSS',10000,'31','Desconto')
        self.cat['rubricas'].append({**c,'cod_inc_cp':'11'})
        self.assertEqual(core.details(self.a)[0]['efeito'],'Pendente')

    def bad_groups(self):
        self.d['bases']={'13':118039}
        self.d['grupos_empresa']=[
            {'base':'13º','grupo':'Alíquota de 20%','base_centavos':4700,'contribuicao_centavos':22667,
             'quantidade_base':1,'aliquota_percentual':20,'pagina':1,'linha_base':'Base empresa sobre 13º Salário 1 47,00',
             'linha_contribuicao':'Valor prev. empresa sobre 13º Salário 2 20% 226,67'},
            {'base':'13º','grupo':'Contribuição zerada','base_centavos':113339,'contribuicao_centavos':0,
             'quantidade_base':2,'aliquota_percentual':None,'pagina':1,'linha_base':'Base empresa sobre 13º Salário 2 1.133,39',
             'linha_contribuicao':'Valor prev. empresa sobre 13º Salário 0 0,00'}]

    def test_inconsistent_rows_suspend_both_group_totals_without_reordering(self):
        self.bad_groups();original=[dict(x) for x in self.d['grupos_empresa']]
        row=core.company_group_rows(self.d)[0]
        self.assertEqual(row['quantidade_contribuicao'],2)
        self.assertEqual(row['validacao'],'A conferir')
        summary=core.company_summary(self.d)[1]
        self.assertIsNone(summary['base_20_centavos'])
        self.assertIsNone(summary['base_zerada_centavos'])
        self.assertEqual(summary['base_nao_determinada_centavos'],118039)
        self.assertEqual(self.d['grupos_empresa'],original)

    def test_manual_decision_and_selection_preserved(self):
        r,c=self.add('0460','13 SAL INDENIZADO',18119)
        self.cat['rubricas'].append({**c,'cod_inc_cp':'00'})
        key=core.key_for(self.d,r)
        self.a['decisions'][key]={'efeito':'Acrescenta','base':'13º','responsavel':'Teste','justificativa':'Memória individual'}
        self.a['selections']={'test|'+key:True}
        row=core.details(self.a)[0]
        self.assertTrue(row['selecionada']);self.assertEqual(row['efeito'],'Acrescenta')
        self.assertEqual(core.reconciliation(self.a)[1]['candidatas'],0)

    def test_value_and_quantity_hint_requires_both(self):
        r,c=self.add('0012','PRO LABORE',3000000,'11',qty='2')
        self.d['grupos_empresa']=[{'base':'Mensal','grupo':'Alíquota de 20%','quantidade_base':2,'aliquota_percentual':20,
                                  'base_centavos':3000000,'contribuicao_centavos':600000,'pagina':1,
                                  'linha_contribuicao':'Valor da previdência empresa 2 20% 6.000,00'}]
        self.assertEqual(len(core.relationship_hints(self.d,core.details(self.a))),1)
        r['quantidade']='1'
        self.assertEqual(core.relationship_hints(self.d,core.details(self.a)),[])

    def test_excel_preserves_raw_values_and_withholds_unreliable_reference(self):
        self.bad_groups();self.add('0290','13 SALARIO',97074)
        w=openpyxl.load_workbook(io.BytesIO(core.export_excel(self.a)))
        s=w['Referencias base empresa'];heads=[x.value for x in s[1]]
        self.assertIsNone(s.cell(3,heads.index('base_20 (R$)')+1).value)
        self.assertIn('Memoria composicao',w.sheetnames)
        s=w['Grupos base empresa'];heads=[x.value for x in s[1]]
        self.assertEqual(s.cell(2,heads.index('base (R$)')+1).value,47)


if __name__=='__main__':unittest.main()
