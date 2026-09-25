import copy
import io
import unittest
import core
import report_layout as layout
import openpyxl
import test_twenty as fixtures

class ReportTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.TwentyTests();self.f.setUp();self.f.mixed()
        self.a=self.f.a;self.d=self.f.d
        self.d['previdencia']={'previdencia_empresa_total':{'valor_centavos':0,'pagina':1}}

    def records(self,w,name):
        s=w[name];values=list(s.values)
        return [dict(zip(values[0],row)) for row in values[1:]]

    def test_overview_preserves_document_grain_and_missing_zero(self):
        before=copy.deepcopy(self.a)
        r=layout.overview(self.a)[0]
        self.assertEqual(len(layout.overview(self.a)),1)
        self.assertEqual(r['base_20_mensal_centavos'],3000000)
        self.assertEqual(r['previdencia_empresa_total_centavos'],0)
        self.assertIsNone(r['rat_total_centavos'])
        self.assertEqual(self.a,before)

    def test_export_has_same_twenty_and_simulation_separately(self):
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        self.assertEqual(w.sheetnames,layout.MAIN_SHEETS)
        t=self.records(w,'Composicao dos 20')[0]
        sim=self.records(w,'Simulacao proporcional')[0]
        self.assertEqual(t[layout.label('parcela_projetada_centavos')],30000)
        expected=core.proportional_twenty(self.a)[0][0]['saldo_estimado_centavos']/100
        self.assertEqual(sim[layout.label('saldo_estimado_centavos')],expected)
        self.assertNotEqual(expected,30000)
        self.assertTrue(w['Cruzamento por folha'].tables)
        self.assertEqual(w['Cruzamento por folha'].freeze_panes,'D2')

    def test_complete_crossing_and_selected_notes_preserve_ids(self):
        self.f.f.add('0009','AJUDA',10000,'00')
        row=core.details(self.a)[1]
        self.a['selections']={core.selection_key(row):True}
        before=copy.deepcopy(self.a)
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        selected=self.records(w,'Selecionadas')[0]
        self.assertEqual(selected['Grupo'],'Fora da base segundo cadastro')
        self.assertEqual(selected['Código'],'0009')
        self.assertEqual(selected['ID do documento'],self.d['hash'])
        self.assertIn('Observação do analista',selected)
        self.assertEqual(len(self.records(w,'Cruzamento por folha')),2)
        self.assertEqual(self.a,before)

    def test_scoped_export_and_formula_like_source_text(self):
        other=copy.deepcopy(self.d);other.update(hash='other',competencia='2014-01')
        self.a['docs'].append(other)
        self.d['rubricas'][0]['descricao']='=1+1'
        a={**self.a,'docs':[self.d]}
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(a)))
        self.assertEqual(len(self.records(w,'Base INSS empresa')),1)
        s=w['Cruzamento por folha'];headers=[c.value for c in s[1]]
        c=s.cell(2,headers.index('Rubrica na folha')+1)
        self.assertEqual(c.value,'=1+1');self.assertEqual(c.data_type,'s')
        self.assertEqual(self.records(w,'Cruzamento por folha')[0]['ID do documento'],self.d['hash'])

    def test_optional_appendices_and_empty_export(self):
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a,include_technical=True)))
        self.assertEqual(w.sheetnames[:len(layout.MAIN_SHEETS)],layout.MAIN_SHEETS)
        self.assertIn('Rubricas detalhadas',w.sheetnames)
        empty={**self.a,'docs':[]}
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(empty)))
        self.assertEqual(w['Base INSS empresa']['A1'].value,'Sem registros neste recorte')

    def test_invalid_reference_stays_blank_and_is_flagged(self):
        self.d['grupos_empresa'][0]['quantidade_base']=3
        w=openpyxl.load_workbook(io.BytesIO(core.export_audit_excel(self.a)))
        r=self.records(w,'Base INSS empresa')[0]
        self.assertIsNone(r[layout.label('base_20_mensal_centavos')])
        self.assertIn('suspensa',r['O que conferir'])
        self.assertEqual(w['Simulacao proporcional']['A1'].value,'Sem registros neste recorte')

if __name__=='__main__':unittest.main()
