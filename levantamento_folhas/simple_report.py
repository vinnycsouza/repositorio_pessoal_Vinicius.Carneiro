"""The same readable payroll blocks for one competence or a consolidated period."""
import io
import math
from datetime import datetime
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SHEETS=('Composição dos 20%', 'Composição da base total', 'Simulação proporcional')
HEADERS=('Código','Rubrica','Valor integral (R$)','Parcela no cenário (R$)',
         'Referência','Critério / observação','Página')

def select_documents(docs,companies,periods,kinds):
    return sorted((d for d in docs if d['cnpj'] in companies and d['competencia'] in periods and d['tipo'] in kinds),
                  key=lambda d:(d['competencia'],d['cnpj'],d['tipo'],d['hash']))

def blocks(a):
    import core
    rows=core.details(a)
    twenty,evidence=core.twenty_composition(a,rows)
    simulations,estimates=core.proportional_twenty(a,rows)
    total=core.reconciliation(a)
    trace=core.composition_trace(a,rows)
    def indexed(records):
        out=defaultdict(list)
        for r in records:out[r['documento']].append(r)
        return out
    indexes=[indexed(x) for x in (rows,twenty,evidence,simulations,estimates,total,trace)]
    result=[]
    for d in sorted(a['docs'],key=lambda d:(d['competencia'],d['cnpj'],d['tipo'],d['hash'])):
        values=[i[d['hash']] for i in indexes]
        result.append(dict(zip(('document','rows','twenty','evidence','simulations','estimates','total','trace'),[d,*values])))
    return result

def export_excel(a):
    import core
    w=openpyxl.Workbook();w.remove(w.active)
    grouped=blocks(a)
    widths=(12,38,20,22,30,66,9)
    money='[$R$-416] #,##0.00;[Red]-[$R$-416] #,##0.00'
    cursor=defaultdict(int)
    def append(s,values,kind=None):
        values=list(values)+[None]*(7-len(values))
        cursor[s.title]+=1
        n=cursor[s.title]
        for col,v in enumerate(values,1):
            c=s.cell(n,col,v)
            if isinstance(v,str):c.data_type='s'
            c.font=Font(name='Arial',size=10,bold=kind in ('title','section','header','total'),
                        color='FFFFFF' if kind in ('title','section','header') else '17324D')
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(v,(int,float)) and col in (3,4):c.number_format=money
            if kind in ('title','section','header'):
                c.fill=PatternFill('solid',fgColor='17324D' if kind=='title' else '93651E' if kind=='section' else '45617B')
            elif kind=='total':c.fill=PatternFill('solid',fgColor='E8EEF4')
        length=max((math.ceil(len(str(v))/(widths[i]*.9)) for i,v in enumerate(values) if v is not None),default=1)
        s.row_dimensions[n].height=min(409,max(25,length*13+7))
        return n
    def text(s,message,kind=None):
        n=append(s,[message],kind)
        s.merge_cells(start_row=n,start_column=1,end_row=n,end_column=7)
        s.row_dimensions[n].height=max(25,math.ceil(len(message)/(sum(widths)*.9))*14+8)
        return n
    def amount(s,label,value,note='',kind='total'):
        n=append(s,[label,None,None,None if value is None else value/100,None,note],kind)
        s.merge_cells(start_row=n,start_column=1,end_row=n,end_column=3)
        if value is None:s.cell(n,4,'Não disponível')
        s.row_dimensions[n].height=max(28,math.ceil(len(note)/(widths[5]*.9))*13+7)
    def rubric(s,r,parcel,criterion=None):
        note=criterion if criterion is not None else r.get('criterio',r.get('motivo',''))
        append(s,[r['codigo'],r['descricao'],r.get('valor_centavos',r.get('valor_integral_centavos'))/100,
                  None if parcel is None else parcel/100,r.get('referencia',''),note,r.get('pagina')])
    def headings(amount_label):
        return (*HEADERS[:3],amount_label,*HEADERS[4:])
    def sections(s,records,field,add_title,reduce_title,amount_label):
        additions=[r for r in records if r['papel']!='Redução indicada']
        reductions=[r for r in records if r['papel']=='Redução indicada']
        for title,items in ((add_title,additions),(reduce_title,reductions)):
            if not items:continue
            text(s,title,'header');append(s,headings(amount_label),'header')
            for r in items:
                # Keep evidence, while removing the duplicate prose already printed in the block.
                criterion=r.get('criterio',r.get('criterio_cadastro',''))
                rubric(s,r,r[field],criterion)
    for name in SHEETS:
        s=w.create_sheet(name)
        for i,width in enumerate(widths,1):s.column_dimensions[openpyxl.utils.get_column_letter(i)].width=width
        s.sheet_view.showGridLines=False;s.sheet_view.zoomScale=85
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
        s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        s.oddFooter.center.text='Página &P de &N'
        text(s,name,'title')
        text(s,f"{a.get('name','Análise')} — {a.get('modelo_relatorio','Consolidado')} | Gerado em {datetime.now():%d/%m/%Y %H:%M} | Versão {core.VERSION}")
        periods=sorted({d['competencia'] for d in a['docs']})
        span=(periods[0]+' a '+periods[-1]) if periods else 'Sem competências'
        text(s,f"Recorte: {len(grouped)} folha(s), {len(periods)} competência(s), de {span}. Filtro de previdência empresa: {a.get('filtro_previdencia','Todas')}.")
        text(s,'Valores em reais. Mensal e 13º permanecem separados. Hipótese e estimativa são cenários independentes: não somar. O Excel é um retrato da análise; alterações não recalculam as classificações.')
        if a.get('catalog',{}).get('identificacao_manual'):
            text(s,'Cadastro de incidência: empresa vinculada manualmente pelo usuário à raiz CNPJ '+a['catalog']['empresa_raiz']+'. O arquivo de incidência não contém identificação da empresa.')
        if not grouped:text(s,'Nenhuma folha no recorte selecionado.')
        if name==SHEETS[2] and grouped and not any(i['simulations'] for i in grouped):
            text(s,'Nenhum grupo misto com referência válida para simulação neste recorte.')
        for item in grouped:
            if name==SHEETS[2] and not item['simulations']:continue
            d=item['document'];rid=d['hash']
            text(s,'')
            text(s,f"{d['competencia']} · {d.get('empresa',d['cnpj'])} · CNPJ {d['cnpj']} · Folha: {d['tipo']}",'title')
            text(s,'Fonte: '+d['arquivo'])
            if d.get('alertas'):text(s,'Alertas do documento: '+'; '.join(d['alertas']))
            # Document-level totals occur once, only on the supporting sheet.
            if name==SHEETS[1]:
                text(s,'Totais do documento — não somar novamente às bases mensal e de 13º','section')
                for r in core.previdencia_rows(d):
                    amount(s,r['indicador'],r['valor_centavos'],f"{r['situacao']}; página {r['pagina'] or 'não localizada'}",kind=None)
                text(s,'Previdência empresa é o valor informado no PDF e não comprova pagamento. Contribuição zerada não comprova desoneração.')
            for base in ('Mensal','13º'):
                t=next(r for r in item['twenty'] if r['base']==base)
                total=next(r for r in item['total'] if r['base']==base)
                parts=[r for r in item['trace'] if r['base']==base]
                sim=next((r for r in item['simulations'] if r['base']==base),None)
                if name==SHEETS[2] and sim is None:continue
                if total['informada_centavos'] is None and not parts:
                    text(s,f'{base}: base não localizada neste documento.');continue
                text(s,'Base '+base,'section')
                if name==SHEETS[0]:
                    amount(s,'Base dos 20% informada no PDF',t['base_20_centavos'],t['situacao'])
                    text(s,t['criterio'])
                    relevant=[r for r in item['evidence'] if r['base']==base]
                    regular=[r for r in relevant if r['parcela_candidata_centavos'] is None]
                    candidates=[r for r in relevant if r['parcela_candidata_centavos'] is not None]
                    sections(s,regular,'parcela_centavos','Rubricas que podem compor a base dos 20%','Reduções da hipótese — detalhamento separado','Parcela na hipótese (R$)')
                    if not regular:text(s,'Sem parcelas atribuídas à hipótese. Consulte a situação da referência acima.')
                    amount(s,'Total projetado / hipótese',t['parcela_projetada_centavos'])
                    amount(s,'Diferença: base dos 20% menos hipótese',t['saldo_nao_identificado_centavos'])
                    if candidates:
                        text(s,'Candidatas conflitantes — fora da hipótese','header');append(s,headings('Candidata (R$)'),'header')
                        for r in candidates:rubric(s,r,r['parcela_candidata_centavos'])
                        amount(s,'Total das candidatas',t['candidatas_conflitantes_centavos'])
                        amount(s,'Diferença se todas forem validadas',t['saldo_condicionado_centavos'],'Fechamento não confirma incidência.')
                    if sim:
                        text(s,'Simulação alternativa disponível na aba Simulação proporcional. Os valores não integram esta hipótese.')
                elif name==SHEETS[2]:
                    text(s,'Estimativa proporcional — cenário independente, a confirmar na folha','section')
                    text(s,sim['premissa'])
                    amount(s,'Base dos 20% de referência',sim['base_20_pdf_centavos'])
                    if t['situacao']=='Hipótese por valor e quantidade':
                        text(s,'Há hipótese específica por valor e quantidade na aba Composição dos 20%. Este rateio é apenas uma alternativa e não a substitui.')
                    text(s,f"Proporção monetária: {core.brl(sim['base_20_pdf_centavos'])} ÷ {core.brl(sim['base_total_pdf_centavos'])} = {sim['fator_proporcao']:.8%}. Cada parcela é arredondada a centavos; não há ajuste para fechar o total.")
                    estimates=[r for r in item['estimates'] if r['base']==base]
                    normal=[r for r in estimates if r['cenario']=='Estimativa pelo cadastro']
                    conditional=[r for r in estimates if r['cenario']!='Estimativa pelo cadastro']
                    sections(s,normal,'estimativa_centavos','Acréscimos estimados','Reduções estimadas','Parcela estimada (R$)')
                    if not normal:text(s,'Sem rubricas elegíveis para estimativa pelo cadastro.')
                    amount(s,'Total de acréscimos estimados',sim['acrescimos_estimados_centavos'])
                    amount(s,'Total de reduções estimadas',sim['reducoes_estimadas_centavos'])
                    amount(s,'Total estimado líquido',sim['saldo_estimado_centavos'],'Acréscimos menos reduções estimados; não somar à hipótese.')
                    amount(s,'Diferença: base dos 20% menos estimativa',sim['diferenca_para_base_20_centavos'])
                    if conditional:
                        text(s,'Estimativa das candidatas conflitantes — fora do total estimado','header');append(s,headings('Candidata estimada (R$)'),'header')
                        for r in conditional:rubric(s,r,r['estimativa_centavos'],r['criterio_cadastro'])
                        amount(s,'Diferença estimada se todas forem validadas',sim['diferenca_com_candidatas_centavos'])
                else:
                    amount(s,'Base total do bloco informada no PDF',total['informada_centavos'],total['situacao_grupos'])
                    regular=[r for r in parts if r['papel'] in ('Acréscimo indicado','Redução indicada')]
                    candidates=[r for r in parts if r['papel']=='Candidata condicionada']
                    sections(s,regular,'parcela_centavos','Possíveis acréscimos à base total','Reduções da base total','Parcela na base total (R$)')
                    if not regular:text(s,'Sem parcelas indicadas pelo cadastro neste bloco.')
                    amount(s,'Total de acréscimos',total['acrescimos_centavos'])
                    amount(s,'Total de reduções',total['reducoes_centavos'])
                    amount(s,'Composição projetada líquida',total['reconstruida_centavos'])
                    amount(s,'Diferença: base total menos composição',total['saldo_sem_explicacao_centavos'],total['conclusao_composicao'])
                    if candidates:
                        text(s,'Candidatas conflitantes — fora da composição projetada','header');append(s,HEADERS,'header')
                        for r in candidates:rubric(s,r,r['parcela_centavos'])
                        amount(s,'Diferença se todas forem validadas',total['saldo_apos_todas_candidatas_centavos'],'Cenário inclui todas as candidatas. Não confirma incidência.')
            if name==SHEETS[1]:
                pending=[r for r in item['trace'] if r['papel']=='Pendente sem valor atribuído']
                if pending:
                    text(s,'Pendências sem parcela atribuída — não somadas às bases','section');append(s,HEADERS,'header')
                    for r in sorted(pending,key=lambda r:-abs(r['valor_centavos'])):rubric(s,r,None)
            text(s,'Diferença negativa indica que a composição excede a base de referência. Valores ausentes não equivalem a zero.')
        s.freeze_panes='C5'
        s.print_options.horizontalCentered=True
        s.print_area=f'A1:G{s.max_row}'
    out=io.BytesIO();w.save(out);return out.getvalue()
