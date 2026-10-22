"""Relatório único da auditoria, sem transformar indícios em créditos."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from auditor_sped_integrado.premissas_icms import PeriodoRegime, UFS
from auditor_sped_integrado.prazo import data_corte


COR_TITULO = "17365D"
COR_CABECALHO = "DCE6F1"
FORMATO_MOEDA = '#,##0.00'
def _valor_celula(valor: object) -> object:
    if pd.isna(valor):
        return None
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        return "'" + valor
    return valor


def resumir_candidatos(c175: pd.DataFrame, simulacao: pd.DataFrame | None = None) -> pd.DataFrame:
    colunas = [
        "competencia", "cnpj_estabelecimento", "fonte_registro", "registros_c175", "valor_operacao",
        "base_pis", "pis_escriturado", "base_cofins", "cofins_escriturada",
        "registros_base_divergente", "registros_chave_invalida",
        "registros_duplicados_entre_abas",
        "contribuicoes_sob_analise", "simulacao_pis_cofins_uf",
        "icms_st_vinculado", "exclusao_calculada", "credito_pis_calculado",
        "credito_cofins_calculado", "estado", "pendencia",
    ]
    if c175.empty:
        return pd.DataFrame(columns=colunas)
    candidatos = c175.loc[
        (c175["cfop"].astype(str) == "5405")
        & (c175["cst_pis"].astype(str).str.zfill(2) == "01")
        & (c175["cst_cofins"].astype(str).str.zfill(2) == "01")
    ]
    if candidatos.empty:
        return pd.DataFrame(columns=colunas)
    if "fonte_registro" not in candidatos.columns:
        candidatos = candidatos.assign(fonte_registro="C175")
    candidatos = candidatos.copy()
    bloqueados = candidatos["registros_duplicados_entre_abas"].fillna(0).gt(0) if "registros_duplicados_entre_abas" in candidatos.columns else pd.Series(False, index=candidatos.index)
    candidatos.loc[bloqueados, ["valor_operacao", "base_pis", "valor_pis", "base_cofins", "valor_cofins"]] = 0
    grupos = candidatos.groupby(["competencia", "cnpj_estabelecimento", "fonte_registro"], as_index=False).agg(
        registros_c175=("registros_c175", "sum"),
        valor_operacao=("valor_operacao", "sum"),
        base_pis=("base_pis", "sum"),
        pis_escriturado=("valor_pis", "sum"),
        base_cofins=("base_cofins", "sum"),
        cofins_escriturada=("valor_cofins", "sum"),
    )
    for coluna in (
        "registros_base_divergente", "registros_chave_invalida",
        "registros_duplicados_entre_abas",
    ):
        if coluna in candidatos.columns:
            contagens = candidatos.groupby(["competencia", "cnpj_estabelecimento", "fonte_registro"])[coluna].sum()
            grupos[coluna] = [
                int(contagens.loc[(linha["competencia"], linha["cnpj_estabelecimento"], linha["fonte_registro"])])
                for _, linha in grupos.iterrows()
            ]
        else:
            grupos[coluna] = 0
    grupos["contribuicoes_sob_analise"] = (
        grupos["pis_escriturado"] + grupos["cofins_escriturada"]
    ).round(2)
    grupos["simulacao_pis_cofins_uf"] = None
    if simulacao is not None and not simulacao.empty:
        if "fonte_registro" not in simulacao.columns:
            simulacao = simulacao.assign(fonte_registro="C175")
        somas = simulacao.groupby(["competencia", "cnpj_estabelecimento", "fonte_registro"])["total_simulado"].sum(min_count=1)
        for indice, linha in grupos.iterrows():
            chave = (linha["competencia"], linha["cnpj_estabelecimento"], linha["fonte_registro"])
            if chave in somas.index and pd.notna(somas.loc[chave]):
                grupos.at[indice, "simulacao_pis_cofins_uf"] = round(float(somas.loc[chave]), 2)
    for coluna in ("icms_st_vinculado", "exclusao_calculada", "credito_pis_calculado", "credito_cofins_calculado"):
        grupos[coluna] = None
    grupos["estado"] = "Possível recuperação em análise"
    grupos["pendencia"] = "Obter NF de aquisição e vincular ICMS-ST à receita para quantificar a exclusão"
    tem_duplicados = grupos["registros_duplicados_entre_abas"] > 0
    grupos.loc[tem_duplicados, "estado"] = "Conferir duplicação entre abas"
    grupos.loc[tem_duplicados, "pendencia"] = "Revisar linhas repetidas entre abas; simulação afetada bloqueada"
    return grupos[colunas]


def gerar_relatorio_excel(
    inventario: list[dict],
    achados: pd.DataFrame,
    st_c170: pd.DataFrame,
    st_c100: pd.DataFrame,
    vendas_c175: pd.DataFrame,
    st_c190: pd.DataFrame,
    regime: str,
    gerado_em: datetime | None = None,
    simulacao: pd.DataFrame | None = None,
    historico_regime: tuple[PeriodoRegime, ...] = (),
    uf_selecionada: str | None = None,
    conferencia_nfe=None,
    data_referencia: date | None = None,
    recorte_prazo: dict[str, int] | None = None,
) -> bytes:
    """Gera três abas: apresentação, memória e validação documental por NF-e."""
    livro = Workbook(write_only=True)
    resumo = livro.create_sheet("Resumo_Executivo")
    memoria = livro.create_sheet("Memoria_Mensal")
    conferencia = livro.create_sheet("Conferencia_NF")

    def adicionar(aba, valores, estilo=None, moedas=(), formulas=()):
        linha = []
        for indice, valor in enumerate(valores, start=1):
            formula_interna = isinstance(valor, str) and valor.startswith("=") and (
                indice in formulas or (aba is conferencia and indice in (21, 22))
            )
            celula = WriteOnlyCell(aba, value=valor if formula_interna else _valor_celula(valor))
            if estilo == "titulo":
                celula.font = Font(size=16, bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor=COR_TITULO)
            elif estilo == "cabecalho":
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor=COR_TITULO)
                celula.alignment = Alignment(wrap_text=True, vertical="center")
            elif estilo == "destaque":
                celula.font = Font(bold=True, color=COR_TITULO)
                celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
            if indice in moedas:
                celula.number_format = FORMATO_MOEDA
            linha.append(celula)
        aba.append(linha)

    candidatos = resumir_candidatos(vendas_c175, simulacao)
    sim = simulacao if simulacao is not None else pd.DataFrame()
    sim_validas = sim.loc[sim["total_simulado"].notna()] if not sim.empty else sim
    anos = sorted(set(candidatos["competencia"].astype(str).str[:4])) if not candidatos.empty else []
    resumo.column_dimensions["A"].width = 35
    for letra in "BCDE":
        resumo.column_dimensions[letra].width = 26
    resumo.column_dimensions["F"].width = 84
    resumo.freeze_panes = "B12"
    resumo.sheet_view.showGridLines = False
    resumo.page_setup.orientation = "landscape"
    resumo.page_setup.fitToWidth = 1
    adicionar(resumo, ["AUDITORIA SPED | ICMS-ST na base de PIS/COFINS", None, None, None, None, None], "titulo")
    adicionar(resumo, ["Análise preliminar para avaliação e conferência documental"])
    adicionar(resumo, ["Gerado em", (gerado_em or datetime.now()).strftime("%d/%m/%Y %H:%M"),
                     "UF escolhida", uf_selecionada or "Pela chave da NF-e"])
    if data_referencia is not None:
        adicionar(resumo, ["Data de referência", data_referencia.strftime("%d/%m/%Y"),
                         "Corte operacional", data_corte(data_referencia).strftime("%d/%m/%Y")])
        adicionar(resumo, ["Recorte de cinco anos", "Competências anteriores e mês que cruza o corte fora dos valores automáticos. Conferir pagamento, medidas judiciais, modulação e aproveitamentos antes de concluir prescrição."])
        contagem = recorte_prazo or {}
        adicionar(resumo, ["Linhas excluídas / pendentes", f"Anteriores: {contagem.get('anterior', 0)}; mês do corte: {contagem.get('limite', 0)}"])
    adicionar(resumo, ["Regime IRPJ/CSLL informado", regime])
    adicionar(resumo, ["Histórico informado", "; ".join(f"{p.inicio} a {p.fim}: {p.regime}" for p in historico_regime) or "Não informado"])
    adicionar(resumo, ["Filtro", "C170 e C175 | CFOP 5405 | CST PIS 01 | CST COFINS 01"])
    adicionar(resumo, ["Escopo", "Possível crédito calculado sobre vendas escrituradas com a alíquota geral da UF. O ICMS-ST da aquisição exige documentos e vínculo às vendas."])
    adicionar(resumo, ["Situação", "Nenhum valor abaixo representa crédito validado ou valor a retificar."])
    adicionar(resumo, ["Alíquotas PIS/COFINS", "Alíquotas escrituradas no registro de origem têm prioridade. Preenchimentos manuais de campos ausentes aparecem identificados na memória e na conferência."])
    adicionar(resumo, [])
    adicionar(resumo, ["INDICADORES", "PIS escriturado sob análise", "COFINS escriturada sob análise", "Possível PIS", "Possível COFINS", "Possível total"], "cabecalho")
    adicionar(resumo, ["Total", float(candidatos["pis_escriturado"].sum()),
                     float(candidatos["cofins_escriturada"].sum()),
                     float(sim_validas["pis_simulado"].sum()) if not sim_validas.empty else None,
                     float(sim_validas["cofins_simulada"].sum()) if not sim_validas.empty else None,
                     float(sim_validas["total_simulado"].sum()) if not sim_validas.empty else None], "destaque", (2, 3, 4, 5, 6))
    adicionar(resumo, [])
    adicionar(resumo, ["ANO", "PIS escriturado", "COFINS escriturada", "Possível PIS", "Possível COFINS", "Possível total"], "cabecalho")
    for ano in anos:
        cand_ano = candidatos.loc[candidatos["competencia"].astype(str).str.startswith(ano)]
        sim_ano = sim_validas.loc[sim_validas["competencia"].astype(str).str.startswith(ano)] if not sim_validas.empty else sim_validas
        adicionar(resumo, [int(ano), float(cand_ano["pis_escriturado"].sum()),
                         float(cand_ano["cofins_escriturada"].sum()),
                         float(sim_ano["pis_simulado"].sum()) if not sim_ano.empty else None,
                         float(sim_ano["cofins_simulada"].sum()) if not sim_ano.empty else None,
                         float(sim_ano["total_simulado"].sum()) if not sim_ano.empty else None], moedas=(2, 3, 4, 5, 6))
    adicionar(resumo, [])
    adicionar(resumo, ["POTENCIAL COM ST DOCUMENTADO", "PIS", "COFINS", "TOTAL"], "cabecalho")
    adicionar(resumo, ["Após conferência manual",
        '=IF(COUNT(Conferencia_NF!U:U)=0,"Pendente",SUM(Conferencia_NF!U:U))',
        '=IF(COUNT(Conferencia_NF!V:V)=0,"Pendente",SUM(Conferencia_NF!V:V))',
        '=IF(COUNT(Conferencia_NF!U:V)=0,"Pendente",SUM(Conferencia_NF!U:V))'],
        "destaque", (2, 3, 4), (2, 3, 4))
    adicionar(resumo, ["Natureza", "Valor sujeito à revisão da apuração, recolhimento, vínculo e aproveitamentos anteriores."])
    adicionar(resumo, [])
    adicionar(resumo, ["COMO LER", "Memoria_Mensal: cálculo por competência, estabelecimento, origem, regime, alíquotas e situação."])
    adicionar(resumo, [None, "Conferencia_NF: uma linha por chave de venda; preencher NF de aquisição e ICMS-ST comprovado alocado à venda."])
    adicionar(resumo, [None, "PIS/COFINS potencial com ST só aparece na linha da NF após conferir documento, ICMS-ST e prazo. Ainda depende de apuração, recolhimento e aproveitamentos anteriores."])
    adicionar(resumo, [])
    adicionar(resumo, ["ARQUIVOS-FONTE", "Nome", "SHA-256"], "cabecalho")
    for item in inventario:
        adicionar(resumo, [item["tipo"], item["arquivo"], item["sha256"]])

    colunas_memoria = ["Competência", "CNPJ estabelecimento", "UF NF-e", "Regime 0110",
                      "Alíquota PIS %", "Alíquota COFINS %", "Registros C175", "Base PIS conferida",
                      "Base COFINS conferida", "Alíquota geral UF %", "Parcela ICMS pela premissa",
                      "Possível PIS", "Possível COFINS", "Possível total", "Bases divergentes",
                      "Chaves inválidas", "Duplicados entre abas", "Situação do cálculo", "Fonte da premissa",
                      "Fonte alíquota PIS", "Fonte alíquota COFINS", "Incidência PIS/COFINS", "Registro de origem"]
    for indice, nome in enumerate(colunas_memoria, start=1):
        memoria.column_dimensions[get_column_letter(indice)].width = min(max(len(nome) + 3, 16), 55)
    memoria.freeze_panes = "C2"
    adicionar(memoria, colunas_memoria, "cabecalho")
    for item in sim.to_dict("records"):
        adicionar(memoria, [item.get(chave) for chave in (
            "competencia", "cnpj_estabelecimento", "uf", "regime_0110", "aliquota_pis",
            "aliquota_cofins", "registros_c175", "base_pis_conferida", "base_cofins_conferida",
            "aliquota_icms_referencia", "parcela_icms_simulada", "pis_simulado", "cofins_simulada",
            "total_simulado", "registros_base_divergente", "registros_chave_invalida",
            "registros_duplicados_entre_abas", "estado_simulacao", "fonte_premissa",
            "fonte_aliquota_pis", "fonte_aliquota_cofins", "descricao_regime", "fonte_registro",
        )], moedas=(8, 9, 11, 12, 13, 14))
    memoria.auto_filter.ref = f"A1:W{max(2, 1 + len(sim))}"

    colunas_nfe = ["Competência", "CNPJ estabelecimento", "Chave NF-e de venda", "UF NF-e",
                   "Incidência PIS/COFINS (0110)", "Itens C175", "Valor operação", "Desconto", "Base PIS",
                   "PIS escriturado", "Base COFINS", "COFINS escriturada", "Alíquota PIS %",
                   "Alíquota COFINS %", "Bases divergentes", "Duplicados entre abas",
                   "Situação automática", "NF(s) de aquisição vinculada(s)", "ICMS-ST comprovado alocado (R$)",
                   "Conferência manual", "PIS potencial com ST", "COFINS potencial com ST",
                   "Revisor", "Observações", "Aba C175", "Primeira linha C175",
                   "Fonte alíquota PIS", "Fonte alíquota COFINS", "Registro de origem", "Sobreposição C170/C175",
                   "Data do pagamento PIS/COFINS", "Medida/ato que afete o prazo", "Prazo conferido"]
    for indice, nome in enumerate(colunas_nfe, start=1):
        conferencia.column_dimensions[get_column_letter(indice)].width = min(max(len(nome) + 3, 14), 52)
    conferencia.column_dimensions["C"].width = 48
    conferencia.column_dimensions["R"].width = 45
    conferencia.column_dimensions["X"].width = 55
    conferencia.freeze_panes = "D2"
    adicionar(conferencia, colunas_nfe, "cabecalho")
    total_nfe = 0
    if conferencia_nfe is not None:
        for total_nfe, linha in enumerate(conferencia_nfe.linhas(), start=1):
            if total_nfe > 1048575:
                raise ValueError("Mais de 1.048.575 NF-e; divida o levantamento por período para caber no Excel.")
            (periodo, cnpj, chave, chave_ok, uf, regime_0110, pis_min, pis_max,
             cof_min, cof_max, itens, oper, desc, base_pis, pis, base_cof, cof,
             divergentes, duplicados, aba, primeira_linha, fonte_registro,
             sobreposicao, fonte_pis, fonte_cofins) = linha
            codigo_regime = regime_0110 or ("1" if (pis_min, cof_min) == (1.65, 7.6) else "2" if (pis_min, cof_min) == (0.65, 3.0) else "")
            pares_aceitos = {
                "1": {(1.65, 7.6)}, "2": {(0.65, 3.0)},
                "3": {(1.65, 7.6), (0.65, 3.0)},
            }
            aliquotas_compativeis = (round(pis_min, 4), round(cof_min, 4)) in pares_aceitos.get(codigo_regime, set())
            situacao = ("Sobreposição C170/C175" if sobreposicao else
                        "Duplicação entre abas/itens" if duplicados else
                        "Chave incompatível" if not chave_ok else
                        "Alíquotas diferentes na NF-e" if pis_min != pis_max or cof_min != cof_max else
                        "Base divergente" if divergentes else
                        "Alíquotas/regime pendentes" if not aliquotas_compativeis else
                        "UF diferente da escolhida" if uf_selecionada and UFS.get(uf) != uf_selecionada else
                        "Apto para conferência")
            excel_linha = total_nfe + 1
            formula_pis = (f'=IF(AND(T{excel_linha}="Confirmado",AG{excel_linha}="Conferido",ISNUMBER(S{excel_linha}),'
                           f'S{excel_linha}>0,Q{excel_linha}="Apto para conferência",R{excel_linha}<>""),'
                           f'MAX(0,MIN(ROUND(MIN(S{excel_linha},I{excel_linha})*M{excel_linha}/100,2),J{excel_linha})),"")')
            formula_cof = (f'=IF(AND(T{excel_linha}="Confirmado",AG{excel_linha}="Conferido",ISNUMBER(S{excel_linha}),'
                           f'S{excel_linha}>0,Q{excel_linha}="Apto para conferência",R{excel_linha}<>""),'
                           f'MAX(0,MIN(ROUND(MIN(S{excel_linha},K{excel_linha})*N{excel_linha}/100,2),L{excel_linha})),"")')
            adicionar(conferencia, [periodo, cnpj, chave, UFS.get(uf, ""), regime_0110,
                itens, oper / 100, desc / 100, base_pis / 100, pis / 100, base_cof / 100,
                cof / 100, pis_min if pis_min == pis_max else None,
                cof_min if cof_min == cof_max else None, divergentes, duplicados,
                situacao, None, None, "Pendente", formula_pis, formula_cof,
                None, None, aba, primeira_linha, fonte_pis or fonte_registro,
                fonte_cofins or fonte_registro, fonte_registro, sobreposicao,
                None, None, "Pendente"],
                moedas=(7, 8, 9, 10, 11, 12, 19, 21, 22))
    conferencia.auto_filter.ref = f"A1:AG{max(2, total_nfe + 1)}"
    validacao = DataValidation(type="list", formula1='"Pendente,Confirmado,Rejeitado"')
    validacao.add(f"T2:T{max(2, total_nfe + 1)}")
    conferencia.data_validations.append(validacao)
    validacao_prazo = DataValidation(type="list", formula1='"Pendente,Conferido,Rejeitado"')
    validacao_prazo.add(f"AG2:AG{max(2, total_nfe + 1)}")
    conferencia.data_validations.append(validacao_prazo)
    arquivo = BytesIO()
    livro.save(arquivo)
    return arquivo.getvalue()
