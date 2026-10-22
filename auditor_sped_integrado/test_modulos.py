from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook

from auditor_sped_integrado.conferencia_nfe import ConferenciaNFe
from auditor_sped_integrado.diagnostico_st import diagnosticar_vendas_st_c170, diagnosticar_vendas_st_c175
from auditor_sped_integrado.modulo_69 import gerar_relatorio_69, levantar_tema_69
from auditor_sped_integrado.premissas_icms import simular_potencial


CNPJ = "10693312000181"
CHAVE_55 = "51230810693312000181550010000000011000000000"
CHAVE_65 = "51230810693312000181650010000000011000000000"


def item_c170(chave: str, *, cfop: int = 5405, item: int = 1, icms: float = 0) -> dict:
    return {
        "Mês": "Agosto", "Ano": 2023, "CNPJ Estabelecimento(C010)": CNPJ,
        "Indicador de Operação(C100)": 1, "Situação(C100)": "00",
        "Chave(C100)": chave, "Modelo(C100)": 55,
        "Numeração Sequencial": item, "CFOP": cfop,
        "Valor Total do Produto": 100, "Valor de Desconto": 0,
        "CST de Pis": "01", "Base de Pis": 100, "Alíquota de Pis": 1.65,
        "Valor de Pis": 1.65, "CST de Cofins": "01", "Base de Cofins": 100,
        "Alíquota de Cofins": 7.6, "Valor de Cofins": 7.6,
        "Valor de Icms": icms, "Valor de Icms ST": 0,
    }


def linha_c175(chave: str, cfop: int = 5405) -> dict:
    return {
        "Mês": "Agosto", "Ano": 2023, "CNPJ Estabelecimento(C010)": CNPJ,
        "Situação(C100)": "00", "Chave(C100)": chave, "Modelo(C100)": 65, "CFOP": cfop,
        "Valor da Operação": 100, "Valor do Desconto": 0,
        "CST PIS": "01", "Valor BC PIS": 100, "Aliquota PIS perc.": 1.65,
        "Valor PIS": 1.65, "CST COFINS": "01", "Valor BC COFINS": 100,
        "Aliquota COFINS perc.": 7.6, "Valor COFINS": 7.6,
    }


class ModulosTest(unittest.TestCase):
    def test_st_inclui_c170_e_c175_sem_duplicar_nfe(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            arquivo = Path(pasta) / "contribuicoes.xlsx"
            with pd.ExcelWriter(arquivo) as livro:
                pd.DataFrame([item_c170(CHAVE_55)]).to_excel(livro, sheet_name="C170 - Itens", index=False)
                pd.DataFrame([linha_c175(CHAVE_65)]).to_excel(livro, sheet_name="C175 - Analítico", index=False)
            with ConferenciaNFe() as notas:
                c170 = diagnosticar_vendas_st_c170(arquivo, coletor_nfe=notas)
                c175 = diagnosticar_vendas_st_c175(arquivo, coletor_nfe=notas)
                self.assertEqual(len(list(notas.linhas())), 2)
                self.assertEqual(notas.sobreposicoes(), set())
            self.assertEqual(c170.iloc[0]["fonte_registro"], "C170")
            self.assertEqual(c175.iloc[0]["fonte_registro"], "C175")
            self.assertEqual(c170.iloc[0]["registros_base_conferida"], 1)
            sim = simular_potencial(pd.concat([c170, c175], ignore_index=True))
            self.assertEqual(set(sim["fonte_registro"]), {"C170", "C175"})

    def test_nfe_presente_em_c170_e_c175_e_sinalizada(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            arquivo = Path(pasta) / "contribuicoes.xlsx"
            with pd.ExcelWriter(arquivo) as livro:
                pd.DataFrame([item_c170(CHAVE_55)]).to_excel(livro, sheet_name="C170", index=False)
                pd.DataFrame([linha_c175(CHAVE_55)]).to_excel(livro, sheet_name="C175", index=False)
            with ConferenciaNFe() as notas:
                diagnosticar_vendas_st_c170(arquivo, coletor_nfe=notas)
                diagnosticar_vendas_st_c175(arquivo, coletor_nfe=notas)
                self.assertEqual(notas.sobreposicoes(), {("2023-08", CNPJ, CHAVE_55)})
                self.assertEqual(list(notas.linhas())[0][22], 1)

    def test_tema_69_separado_usa_icms_escriturado_e_bloqueia_duplicacao(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            icms = Path(pasta) / "icms.xlsx"
            pis = Path(pasta) / "pis.xlsx"
            venda = item_c170(CHAVE_55, cfop=5102, icms=18)
            pd.DataFrame([venda]).rename(columns={
                "CNPJ Estabelecimento(C010)": "CNPJ",
                "Chave(C100)": "Chave de Acesso(C100)",
            }).to_excel(icms, sheet_name="C170", index=False)
            pd.DataFrame([venda]).to_excel(pis, sheet_name="C170", index=False)
            resultado, contagens = levantar_tema_69(icms, pis)
            self.assertEqual(len(resultado), 1)
            self.assertAlmostEqual(resultado.iloc[0]["pis_potencial"], 0.30)
            self.assertAlmostEqual(resultado.iloc[0]["cofins_potencial"], 1.37)
            self.assertEqual(contagens["itens_com_calculo_preliminar"], 1)
            relatorio = load_workbook(BytesIO(gerar_relatorio_69([], resultado, contagens)))
            self.assertEqual(relatorio.sheetnames, ["Resumo_Tema_69", "Memoria_Mensal_69", "Conferencia_Tema_69"])
            self.assertEqual(relatorio["Conferencia_Tema_69"].cell(2, 3).value, CHAVE_55)
            self.assertEqual(relatorio["Memoria_Mensal_69"].cell(2, 13).value, 1.67)
            resumo_69 = relatorio["Resumo_Tema_69"]
            total_69 = next(linha for linha in resumo_69 if linha[0].value == "Possível total sob conferência")
            self.assertEqual(total_69[1].value, 1.67)
            notas_69 = relatorio["Conferencia_Tema_69"]
            self.assertEqual(notas_69.freeze_panes, "D2")
            self.assertEqual(notas_69["R2"].value, "Pendente")
            self.assertEqual(notas_69["S2"].value, "Pendente")
            self.assertEqual(notas_69["R2"].fill.fgColor.rgb[-6:], "FFF2CC")
            self.assertEqual(len(notas_69.data_validations.dataValidation), 2)
            self.assertEqual(len(notas_69.conditional_formatting), 2)
            relatorio.close()

            with pd.ExcelWriter(pis) as livro:
                pd.DataFrame([venda]).to_excel(livro, sheet_name="C170", index=False)
                pd.DataFrame([venda]).to_excel(livro, sheet_name="C170--2", index=False)
            duplicado, _ = levantar_tema_69(icms, pis)
            self.assertTrue(duplicado["total_potencial"].isna().all())
            self.assertIn("duplicado", duplicado.iloc[0]["motivos"])

    def test_tema_69_encontra_c190_quando_icms_nao_tem_saida_c170(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            icms = Path(pasta) / "icms.xlsx"
            pis = Path(pasta) / "pis.xlsx"
            c190 = {
                "Mês": "Agosto", "Ano": 2023, "CNPJ": CNPJ,
                "Indicador de Operação(C100)": 1, "Chave de Acesso(C100)": CHAVE_55,
                "CFOP": 5102, "CST de ICMS": "000", "Alíquota de Icms": 18,
                "Base de Icms": 100, "Valor de Icms": 18,
            }
            pd.DataFrame([c190]).to_excel(icms, sheet_name="C190", index=False)
            pd.DataFrame([item_c170(CHAVE_55, cfop=5102)]).to_excel(pis, sheet_name="C170", index=False)
            resultado, _ = levantar_tema_69(icms, pis)
            self.assertEqual(resultado.iloc[0]["fonte_icms"], "C190")
            self.assertAlmostEqual(resultado.iloc[0]["total_potencial"], 1.67)

    def test_tema_69_concilia_c175_e_c190_sem_somar_documento_duas_vezes(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            icms = Path(pasta) / "icms.xlsx"
            pis = Path(pasta) / "pis.xlsx"
            c190 = {
                "Mês": "Agosto", "Ano": 2023, "CNPJ": CNPJ,
                "Indicador de Operação(C100)": 1, "Chave de Acesso(C100)": CHAVE_65,
                "CFOP": 5102, "CST de ICMS": "000", "Alíquota de Icms": 18,
                "Base de Icms": 100, "Valor de Icms": 18,
            }
            pd.DataFrame([c190]).to_excel(icms, sheet_name="C190", index=False)
            pd.DataFrame([linha_c175(CHAVE_65, 5102)]).to_excel(pis, sheet_name="C175", index=False)
            resultado, _ = levantar_tema_69(icms, pis)
            self.assertEqual(len(resultado), 1)
            self.assertEqual(resultado.iloc[0]["fonte_receita"], "C175")
            self.assertAlmostEqual(resultado.iloc[0]["total_potencial"], 1.67)

            with pd.ExcelWriter(pis) as livro:
                pd.DataFrame([linha_c175(CHAVE_65, 5102)]).to_excel(livro, sheet_name="C175", index=False)
                pd.DataFrame([linha_c175(CHAVE_65, 5102)]).to_excel(livro, sheet_name="C175--2", index=False)
            duplicado, _ = levantar_tema_69(icms, pis)
            self.assertTrue(duplicado["total_potencial"].isna().all())
            self.assertIn("duplicado", duplicado.iloc[0]["motivos"])

    def test_tema_69_bloqueia_c190_quando_cfop_mistura_cst_tributado_e_nao_tributado(self):
        with TemporaryDirectory(dir=Path(__file__).resolve().parent) as pasta:
            icms = Path(pasta) / "icms.xlsx"
            pis = Path(pasta) / "pis.xlsx"
            c190 = {
                "Mês": "Agosto", "Ano": 2023, "CNPJ": CNPJ,
                "Indicador de Operação(C100)": 1, "Chave de Acesso(C100)": CHAVE_65,
                "CFOP": 5102, "CST de ICMS": "000", "Alíquota de Icms": 18,
                "Base de Icms": 200, "Valor de Icms": 36,
            }
            pd.DataFrame([c190]).to_excel(icms, sheet_name="C190", index=False)
            tributado = linha_c175(CHAVE_65, 5102)
            nao_tributado = {**tributado, "CST PIS": "04", "CST COFINS": "04",
                             "Valor PIS": 0, "Valor COFINS": 0}
            pd.DataFrame([tributado, nao_tributado]).to_excel(pis, sheet_name="C175", index=False)
            resultado, _ = levantar_tema_69(icms, pis)
            self.assertTrue(resultado["total_potencial"].isna().all())
            self.assertIn("CST não tributado", resultado.iloc[0]["motivos"])


if __name__ == "__main__":
    unittest.main()
