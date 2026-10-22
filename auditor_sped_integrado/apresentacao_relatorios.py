"""Convenções compartilhadas de cobertura e apresentação dos relatórios."""

from __future__ import annotations

from datetime import date

from openpyxl.styles import Alignment, Font, PatternFill

from auditor_sped_integrado.prazo import data_corte


AZUL_ESCURO = "17365D"
AZUL_CLARO = "EAF2F8"
AMBAR_CLARO = "FFF2CC"
CINZA_CLARO = "EEF1F4"
VERMELHO_CLARO = "FCE8E6"
VERDE_CLARO = "E2F0D9"
BRANCO = "FFFFFF"

PREENCHIMENTOS = {
    "titulo": PatternFill("solid", fgColor=AZUL_ESCURO),
    "cabecalho": PatternFill("solid", fgColor=AZUL_ESCURO),
    "entrada": PatternFill("solid", fgColor=AMBAR_CLARO),
    "calculo": PatternFill("solid", fgColor=AZUL_CLARO),
    "tecnico": PatternFill("solid", fgColor=CINZA_CLARO),
    "pendencia": PatternFill("solid", fgColor=VERMELHO_CLARO),
    "confirmado": PatternFill("solid", fgColor=VERDE_CLARO),
}


def estilizar_cabecalho(celula, grupo: str = "cabecalho") -> None:
    celula.font = Font(bold=True, color=BRANCO if grupo in {"cabecalho", "titulo"} else AZUL_ESCURO)
    celula.fill = PREENCHIMENTOS[grupo]
    celula.alignment = Alignment(wrap_text=True, vertical="center")


def mes_seguinte(competencia: str) -> str:
    ano, mes = map(int, competencia.split("-"))
    return f"{ano + (mes == 12):04d}-{1 if mes == 12 else mes + 1:02d}"


def competencias_esperadas(referencia: date | None, observadas) -> list[str]:
    """Meses completos do recorte; ausência significa apenas nenhum grupo no relatório."""
    observadas = sorted({str(valor)[:7] for valor in observadas if str(valor)[:7].count("-") == 1})
    if referencia is not None:
        primeiro = mes_seguinte(data_corte(referencia).strftime("%Y-%m"))
        ultimo = (referencia.replace(day=1).toordinal() - 1)
        ultimo_mes = date.fromordinal(ultimo).strftime("%Y-%m")
    elif observadas:
        primeiro, ultimo_mes = observadas[0], observadas[-1]
    else:
        return []
    meses = []
    atual = primeiro
    while atual <= ultimo_mes:
        meses.append(atual)
        atual = mes_seguinte(atual)
    return meses


def lacunas_competencia(referencia: date | None, observadas) -> list[str]:
    conjunto = {str(valor)[:7] for valor in observadas}
    return [mes for mes in competencias_esperadas(referencia, conjunto) if mes not in conjunto]


def cobertura_do_mes(mes: str, tem_grupo: bool, fontes: dict | None) -> tuple[int | None, int | None, str]:
    """Expõe presença de saídas mapeadas sem inferir arquivo ausente ou prazo fiscal."""
    if not fontes:
        return None, None, "Com grupo no filtro" if tem_grupo else "Sem grupo no filtro; cobertura das fontes não informada"
    pis = int(fontes.get("EFD-Contribuições", {}).get(mes, 0))
    icms = int(fontes.get("EFD ICMS/IPI", {}).get(mes, 0)) if fontes.get("icms_fornecido") else None
    if tem_grupo:
        situacao = "Com grupo no filtro"
    elif pis and icms:
        situacao = "Saídas nas duas fontes; sem grupo no filtro"
    elif pis:
        situacao = "Saídas em Contribuições; sem grupo no filtro; conferir ICMS/IPI"
    elif icms:
        situacao = "Saídas C190 no ICMS/IPI; sem saída mapeada em Contribuições"
    else:
        situacao = "Sem saída mapeada nas fontes; verificar cobertura dos arquivos"
    return pis, icms, situacao
