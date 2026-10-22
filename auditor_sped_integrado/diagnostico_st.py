"""Quantifica ICMS-ST destacado no C170 da EFD-Contribuições do Sysconv.

O resultado é evidência de escrituração, não valor de crédito tributário.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from hashlib import blake2b
from io import BytesIO
from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from auditor_sped_integrado.motor import _cnpj, _competencia


COLUNAS = (
    "competencia", "cnpj_estabelecimento", "cfop", "tipo_operacao", "regime_0110",
    "itens_com_st", "valor_icms_st_destacado", "primeira_aba_c170", "primeira_linha_c170",
)
COLUNAS_C100 = (
    "aba", "competencia", "cnpj_estabelecimento", "indicador_operacao",
    "documentos_com_st", "valor_icms_st_destacado", "primeira_linha_c100",
)
COLUNAS_C175 = (
    "competencia", "cnpj_estabelecimento", "uf_nfe", "chave_conferida",
    "regime_0110", "cfop", "cst_pis", "cst_cofins",
    "aliquota_pis", "aliquota_cofins", "fonte_aliquota_pis", "fonte_aliquota_cofins",
    "registros_c175", "valor_operacao",
    "valor_desconto", "base_pis", "valor_pis", "base_cofins", "valor_cofins",
    "registros_base_conferida", "base_pis_conferida", "base_cofins_conferida",
    "valor_pis_conferido", "valor_cofins_conferido", "registros_base_divergente",
    "registros_chave_invalida", "registros_duplicados_entre_abas",
    "primeira_aba", "primeira_linha_c175", "fonte_registro",
)
COLUNAS_C190 = (
    "competencia", "cnpj_estabelecimento", "cfop", "indicador_operacao",
    "registros_c190", "valor_icms_st_destacado", "primeira_linha_c190",
)
TRANSFERENCIAS = {"1152", "1409", "2152", "2409", "5152", "5409", "6152", "6409"}
CODIGOS_UF = {
    "11", "12", "13", "14", "15", "16", "17", "21", "22", "23", "24", "25",
    "26", "27", "28", "29", "31", "32", "33", "35", "41", "42", "43",
    "50", "51", "52", "53",
}


def _slug(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", texto).strip().casefold()


def _indice(cabecalho: tuple, *aliases: str) -> int:
    campos = {_slug(valor): indice for indice, valor in enumerate(cabecalho)}
    for alias in aliases:
        if _slug(alias) in campos:
            return campos[_slug(alias)]
    raise ValueError(f"Coluna obrigatória ausente no C170: {aliases[0]}")


def _decimal(valor: object) -> Decimal:
    if valor is None or str(valor).strip() == "":
        return Decimal("0")
    try:
        return Decimal(str(valor).replace(".", "").replace(",", ".") if isinstance(valor, str) and "," in valor else str(valor))
    except InvalidOperation as erro:
        raise ValueError(f"Valor de ICMS-ST inválido: {valor}") from erro


def _regimes_0110(livro) -> dict[str, str]:
    if "0110" not in livro.sheetnames:
        return {}
    aba = livro["0110"]
    aba.reset_dimensions()
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas)
    mes_i, ano_i = _indice(cabecalho, "Mês"), _indice(cabecalho, "Ano")
    regime_i = _indice(cabecalho, "COD_INC_TRIB")
    regimes = {}
    for linha in linhas:
        periodo = _competencia(linha[mes_i], linha[ano_i])
        if not periodo:
            continue
        codigo = str(linha[regime_i] or "").strip().removesuffix(".0")
        if codigo not in {"1", "2", "3"}:
            raise ValueError(f"Regime 0110 inválido em {periodo}: {codigo}")
        if periodo in regimes and regimes[periodo] != codigo:
            raise ValueError(f"Regimes 0110 conflitantes em {periodo}.")
        regimes[periodo] = codigo
    return regimes


def diagnosticar_st_c170(arquivo: str | Path | bytes) -> pd.DataFrame:
    """Soma ICMS-ST positivo em entradas C170, sem inferir direito a crédito."""
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        regimes = _regimes_0110(livro)
        abas_c170 = [nome for nome in livro.sheetnames if _slug(nome).startswith("c170")]
        if not abas_c170:
            raise ValueError("Aba C170 não encontrada na EFD-Contribuições.")
        grupos = defaultdict(lambda: {"itens": 0, "valor": Decimal("0"), "primeira_aba": None, "primeira_linha": None})
        for nome_c170 in abas_c170:
            aba = livro[nome_c170]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas)
            mes_i, ano_i = _indice(cabecalho, "Mês"), _indice(cabecalho, "Ano")
            cnpj_i = _indice(cabecalho, "CNPJ Estabelecimento(C010)", "CNPJ Estabelecimento", "CNPJ")
            oper_i = _indice(cabecalho, "Indicador de Operação(C100)", "Indicador de Operação")
            try:
                situacao_i = _indice(cabecalho, "Situação(C100)", "Situação")
            except ValueError:
                situacao_i = None
            cfop_i = _indice(cabecalho, "CFOP")
            st_i = _indice(cabecalho, "Valor de Icms ST", "Valor de ICMS-ST", "VL_ICMS_ST")
            for numero_linha, linha in enumerate(linhas, start=2):
                if str(linha[oper_i] if linha[oper_i] is not None else "").strip() not in {"0", "0.0"}:
                    continue
                if situacao_i is not None and str(linha[situacao_i] if linha[situacao_i] is not None else "").strip().zfill(2) not in {"00", "01"}:
                    continue
                valor_st = _decimal(linha[st_i])
                if valor_st <= 0:
                    continue
                periodo = _competencia(linha[mes_i], linha[ano_i])
                cnpj = _cnpj(linha[cnpj_i])
                cfop = re.sub(r"\D", "", str(linha[cfop_i] or ""))
                if not periodo or not cnpj or len(cfop) != 4:
                    raise ValueError(f"Identificação incompleta em {nome_c170}, linha {numero_linha}.")
                tipo = "Transferência" if cfop in TRANSFERENCIAS else "Outra entrada"
                grupo = grupos[(periodo, cnpj, cfop, tipo, regimes.get(periodo, ""))]
                grupo["itens"] += 1
                grupo["valor"] += valor_st
                if grupo["primeira_linha"] is None:
                    grupo["primeira_aba"] = nome_c170
                    grupo["primeira_linha"] = numero_linha
        saida = [
            {
                "competencia": periodo, "cnpj_estabelecimento": cnpj, "cfop": cfop,
                "tipo_operacao": tipo, "regime_0110": regime,
                "itens_com_st": dados["itens"],
                "valor_icms_st_destacado": float(dados["valor"].quantize(Decimal("0.01"))),
                "primeira_aba_c170": dados["primeira_aba"],
                "primeira_linha_c170": dados["primeira_linha"],
            }
            for (periodo, cnpj, cfop, tipo, regime), dados in grupos.items()
        ]
        return pd.DataFrame(saida, columns=COLUNAS).sort_values(
            ["competencia", "cnpj_estabelecimento", "cfop"], ignore_index=True
        ) if saida else pd.DataFrame(columns=COLUNAS)
    finally:
        livro.close()


def diagnosticar_st_c100(arquivo: str | Path | bytes) -> pd.DataFrame:
    """Localiza o ST destacado no documento, inclusive em abas C100 divididas."""
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        abas = [nome for nome in livro.sheetnames if _slug(nome).startswith("c100")]
        if not abas:
            raise ValueError("Aba C100 não encontrada na EFD-Contribuições.")
        grupos = defaultdict(lambda: {"documentos": 0, "valor": Decimal("0"), "primeira_linha": None})
        for nome in abas:
            aba = livro[nome]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas)
            mes_i, ano_i = _indice(cabecalho, "Mês"), _indice(cabecalho, "Ano")
            cnpj_i = _indice(cabecalho, "CNPJ Estabelecimento(C010)", "CNPJ Estabelecimento", "CNPJ")
            oper_i = _indice(cabecalho, "Indicador de Operação")
            situacao_i = _indice(cabecalho, "Situação")
            st_i = _indice(cabecalho, "Valor de ICMS - ST", "Valor de ICMS ST", "VL_ICMS_ST")
            for numero_linha, linha in enumerate(linhas, start=2):
                if str(linha[situacao_i] if linha[situacao_i] is not None else "").strip().zfill(2) not in {"00", "01"}:
                    continue
                valor_st = _decimal(linha[st_i])
                if valor_st <= 0:
                    continue
                periodo = _competencia(linha[mes_i], linha[ano_i])
                cnpj = _cnpj(linha[cnpj_i])
                operacao = str(linha[oper_i] if linha[oper_i] is not None else "").strip().removesuffix(".0")
                if not periodo or not cnpj or operacao not in {"0", "1"}:
                    raise ValueError(f"Identificação incompleta em {nome}, linha {numero_linha}.")
                grupo = grupos[(nome, periodo, cnpj, operacao)]
                grupo["documentos"] += 1
                grupo["valor"] += valor_st
                if grupo["primeira_linha"] is None:
                    grupo["primeira_linha"] = numero_linha
        saida = [
            {
                "aba": nome, "competencia": periodo, "cnpj_estabelecimento": cnpj,
                "indicador_operacao": operacao, "documentos_com_st": dados["documentos"],
                "valor_icms_st_destacado": float(dados["valor"].quantize(Decimal("0.01"))),
                "primeira_linha_c100": dados["primeira_linha"],
            }
            for (nome, periodo, cnpj, operacao), dados in grupos.items()
        ]
        return pd.DataFrame(saida, columns=COLUNAS_C100).sort_values(
            ["competencia", "cnpj_estabelecimento", "aba", "indicador_operacao"], ignore_index=True
        ) if saida else pd.DataFrame(columns=COLUNAS_C100)
    finally:
        livro.close()


def diagnosticar_vendas_st_c175(
    arquivo: str | Path | bytes, cfops: set[str] | None = None, coletor_nfe=None,
    aliquotas_informadas: dict[tuple[str, str, float | None, float | None], tuple[float, float]] | None = None,
    data_referencia=None,
) -> pd.DataFrame:
    """Resume vendas com CFOP de ST no C175, sem atribuir ICMS-ST não informado."""
    cfops = cfops or {"5405"}
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        abas = [nome for nome in livro.sheetnames if _slug(nome).startswith("c175")]
        if not abas:
            raise ValueError("Aba C175 não encontrada na EFD-Contribuições.")
        regimes = _regimes_0110(livro)
        from auditor_sped_integrado.prazo import classificar_competencia
        mapa_cfop = Counter()
        cobertura_mensal = Counter()
        recorte = Counter()
        assinaturas_primeira_aba = {}
        grupos = defaultdict(lambda: {"registros": 0, "operacao": Decimal("0"), "base_pis": Decimal("0"),
                                     "pis": Decimal("0"), "base_cofins": Decimal("0"), "cofins": Decimal("0"),
                                     "desconto": Decimal("0"), "conferidos": 0, "divergentes": 0,
                                     "chaves_invalidas": 0, "duplicados": 0,
                                     "base_pis_ok": Decimal("0"),
                                     "base_cofins_ok": Decimal("0"), "pis_ok": Decimal("0"),
                                     "cofins_ok": Decimal("0"),
                                     "primeira_aba": None, "primeira_linha": None})
        for nome in abas:
            aba = livro[nome]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas)
            mes_i, ano_i = _indice(cabecalho, "Mês"), _indice(cabecalho, "Ano")
            cnpj_i = _indice(cabecalho, "CNPJ Estabelecimento(C010)", "CNPJ Estabelecimento", "CNPJ")
            cfop_i = _indice(cabecalho, "CFOP")
            situacao_i = _indice(cabecalho, "Situação(C100)", "Situação")
            cst_pis_i, cst_cofins_i = _indice(cabecalho, "CST PIS"), _indice(cabecalho, "CST COFINS")
            aliquota_pis_i = _indice(cabecalho, "Aliquota PIS perc.")
            aliquota_cofins_i = _indice(cabecalho, "Aliquota COFINS perc.")
            operacao_i = _indice(cabecalho, "Valor da Operação")
            try:
                desconto_i = _indice(cabecalho, "Valor do Desconto")
            except ValueError:
                desconto_i = None
            try:
                chave_i = _indice(cabecalho, "Chave(C100)", "Chave", "CHV_NFE")
            except ValueError:
                chave_i = None
            base_pis_i, pis_i = _indice(cabecalho, "Valor BC PIS"), _indice(cabecalho, "Valor PIS")
            base_cofins_i, cofins_i = _indice(cabecalho, "Valor BC COFINS"), _indice(cabecalho, "Valor COFINS")
            for numero_linha, linha in enumerate(linhas, start=2):
                periodo = _competencia(linha[mes_i], linha[ano_i])
                if data_referencia is not None:
                    estado_prazo = classificar_competencia(periodo, data_referencia)
                    recorte[estado_prazo] += 1
                    if estado_prazo != "dentro":
                        continue
                cfop = re.sub(r"\D", "", str(linha[cfop_i] or ""))
                if len(cfop) == 4 and cfop[0] in {"5", "6"} and str(linha[situacao_i] or "").strip().removesuffix(".0").zfill(2) in {"00", "01"}:
                    mapa_cfop[("C175", cfop, str(linha[cst_pis_i] or "").strip().removesuffix(".0").zfill(2),
                               str(linha[cst_cofins_i] or "").strip().removesuffix(".0").zfill(2))] += 1
                    cobertura_mensal[periodo] += 1
                if cfop not in cfops:
                    continue
                if str(linha[situacao_i] if linha[situacao_i] is not None else "").strip().zfill(2) not in {"00", "01"}:
                    continue
                periodo = _competencia(linha[mes_i], linha[ano_i])
                cnpj = _cnpj(linha[cnpj_i])
                if not periodo or not cnpj:
                    raise ValueError(f"Identificação incompleta em {nome}, linha {numero_linha}.")
                cst_pis = str(linha[cst_pis_i] or "").strip().zfill(2)
                cst_cofins = str(linha[cst_cofins_i] or "").strip().zfill(2)
                pis_presente = linha[aliquota_pis_i] is not None and str(linha[aliquota_pis_i]).strip() != ""
                cofins_presente = linha[aliquota_cofins_i] is not None and str(linha[aliquota_cofins_i]).strip() != ""
                aliquota_pis = _decimal(linha[aliquota_pis_i])
                aliquota_cofins = _decimal(linha[aliquota_cofins_i])
                if cst_pis == "01" and cst_cofins == "01" and cfop == "5405" and (not pis_presente or not cofins_presente):
                    chave_escolha = (
                        periodo, cnpj,
                        float(aliquota_pis) if pis_presente else None,
                        float(aliquota_cofins) if cofins_presente else None,
                    )
                    escolha = (aliquotas_informadas or {}).get(chave_escolha)
                    if escolha is not None:
                        if not pis_presente:
                            aliquota_pis = Decimal(str(escolha[0]))
                        if not cofins_presente:
                            aliquota_cofins = Decimal(str(escolha[1]))
                        if aliquota_pis <= 0 or aliquota_cofins <= 0:
                            raise ValueError(f"Alíquota informada inválida para {periodo}, estabelecimento {cnpj}.")
                fonte_pis = "C175" if pis_presente else "Informada pelo usuário" if aliquota_pis > 0 else "Ausente"
                fonte_cofins = "C175" if cofins_presente else "Informada pelo usuário" if aliquota_cofins > 0 else "Ausente"
                chave = re.sub(r"\D", "", str(linha[chave_i] or "")) if chave_i is not None else ""
                chave_valida = (
                    len(chave) == 44 and chave[:2] in CODIGOS_UF and chave[6:20] == cnpj
                    and chave[2:6] == periodo[2:4] + periodo[5:7]
                )
                # A UF da chave é preservada mesmo quando a validação completa falha.
                uf_nfe = chave[:2] if len(chave) == 44 else ""
                grupo = grupos[(periodo, cnpj, uf_nfe, chave_valida, regimes.get(periodo, ""),
                                 cfop, cst_pis, cst_cofins, aliquota_pis, aliquota_cofins, fonte_pis, fonte_cofins)]
                if cst_pis == "01" and cst_cofins == "01" and len(abas) > 1:
                    assinatura = blake2b(repr(tuple(linha)).encode("utf-8"), digest_size=16).digest()
                    primeira_aba = assinaturas_primeira_aba.setdefault(assinatura, nome)
                    if primeira_aba != nome:
                        grupo["duplicados"] += 1
                        if coletor_nfe is not None and cfop == "5405":
                            coletor_nfe.adicionar(
                                competencia=periodo, cnpj=cnpj, chave=chave, chave_ok=chave_valida,
                                uf=uf_nfe, regime=regimes.get(periodo, ""),
                                aliquota_pis=aliquota_pis, aliquota_cofins=aliquota_cofins,
                                operacao=Decimal("0"), desconto=Decimal("0"),
                                base_pis=Decimal("0"), pis=Decimal("0"),
                                base_cofins=Decimal("0"), cofins=Decimal("0"),
                                base_ok=True, duplicado=True, aba=nome, linha=numero_linha,
                                fonte_pis=fonte_pis, fonte_cofins=fonte_cofins,
                            )
                        continue
                grupo["registros"] += 1
                operacao = _decimal(linha[operacao_i])
                desconto = _decimal(linha[desconto_i]) if desconto_i is not None else Decimal("0")
                base_pis = _decimal(linha[base_pis_i])
                pis = _decimal(linha[pis_i])
                base_cofins = _decimal(linha[base_cofins_i])
                cofins = _decimal(linha[cofins_i])
                grupo["operacao"] += operacao
                grupo["desconto"] += desconto
                grupo["base_pis"] += base_pis
                grupo["pis"] += pis
                grupo["base_cofins"] += base_cofins
                grupo["cofins"] += cofins
                base_ok = (
                    desconto_i is not None
                    and abs(operacao - desconto - base_pis) <= Decimal("0.01")
                    and abs(base_pis - base_cofins) <= Decimal("0.01")
                    and (cst_pis != "01" or abs(base_pis * aliquota_pis / Decimal(100) - pis) <= Decimal("0.02"))
                    and (cst_cofins != "01" or abs(base_cofins * aliquota_cofins / Decimal(100) - cofins) <= Decimal("0.02"))
                )
                if base_ok:
                    grupo["conferidos"] += 1
                    grupo["base_pis_ok"] += base_pis
                    grupo["base_cofins_ok"] += base_cofins
                    grupo["pis_ok"] += pis
                    grupo["cofins_ok"] += cofins
                else:
                    grupo["divergentes"] += 1
                if not chave_valida:
                    grupo["chaves_invalidas"] += 1
                if coletor_nfe is not None and cst_pis == "01" and cst_cofins == "01" and cfop == "5405":
                    coletor_nfe.adicionar(
                        competencia=periodo, cnpj=cnpj, chave=chave, chave_ok=chave_valida,
                        uf=uf_nfe, regime=regimes.get(periodo, ""),
                        aliquota_pis=aliquota_pis, aliquota_cofins=aliquota_cofins,
                        operacao=operacao, desconto=desconto, base_pis=base_pis, pis=pis,
                        base_cofins=base_cofins, cofins=cofins,
                        base_ok=base_ok, duplicado=False, aba=nome, linha=numero_linha,
                        fonte_pis=fonte_pis, fonte_cofins=fonte_cofins,
                    )
                if grupo["primeira_linha"] is None:
                    grupo["primeira_aba"], grupo["primeira_linha"] = nome, numero_linha
        saida = [
            {
                "competencia": periodo, "cnpj_estabelecimento": cnpj, "uf_nfe": uf_nfe,
                "chave_conferida": chave_valida,
                "regime_0110": regime, "cfop": cfop,
                "cst_pis": cst_pis, "cst_cofins": cst_cofins,
                "aliquota_pis": float(aliquota_pis), "aliquota_cofins": float(aliquota_cofins),
                "fonte_aliquota_pis": fonte_pis, "fonte_aliquota_cofins": fonte_cofins,
                "registros_c175": dados["registros"],
                "valor_operacao": float(dados["operacao"].quantize(Decimal("0.01"))),
                "valor_desconto": float(dados["desconto"].quantize(Decimal("0.01"))),
                "base_pis": float(dados["base_pis"].quantize(Decimal("0.01"))),
                "valor_pis": float(dados["pis"].quantize(Decimal("0.01"))),
                "base_cofins": float(dados["base_cofins"].quantize(Decimal("0.01"))),
                "valor_cofins": float(dados["cofins"].quantize(Decimal("0.01"))),
                "registros_base_conferida": dados["conferidos"],
                "base_pis_conferida": float(dados["base_pis_ok"].quantize(Decimal("0.01"))),
                "base_cofins_conferida": float(dados["base_cofins_ok"].quantize(Decimal("0.01"))),
                "valor_pis_conferido": float(dados["pis_ok"].quantize(Decimal("0.01"))),
                "valor_cofins_conferido": float(dados["cofins_ok"].quantize(Decimal("0.01"))),
                "registros_base_divergente": dados["divergentes"],
                "registros_chave_invalida": dados["chaves_invalidas"],
                "registros_duplicados_entre_abas": dados["duplicados"],
                "primeira_aba": dados["primeira_aba"], "primeira_linha_c175": dados["primeira_linha"],
                "fonte_registro": "C175",
            }
            for (periodo, cnpj, uf_nfe, chave_valida, regime, cfop, cst_pis, cst_cofins, aliquota_pis, aliquota_cofins, fonte_pis, fonte_cofins), dados in grupos.items()
        ]
        resultado = pd.DataFrame(saida, columns=COLUNAS_C175).sort_values(
            ["competencia", "cnpj_estabelecimento", "cfop", "cst_pis", "cst_cofins"], ignore_index=True
        ) if saida else pd.DataFrame(columns=COLUNAS_C175)
        resultado.attrs["mapa_cfop"] = mapa_cfop
        resultado.attrs["cobertura_mensal"] = dict(cobertura_mensal)
        resultado.attrs["recorte_prazo"] = dict(recorte)
        return resultado
    finally:
        livro.close()


def diagnosticar_vendas_st_c170(
    arquivo: str | Path | bytes, coletor_nfe=None,
    aliquotas_informadas: dict | None = None,
    data_referencia=None,
) -> pd.DataFrame:
    """Lê saídas C170 tributadas, preservando item e origem para conferência.

    A base de item pode incluir frete e outras parcelas; não se exige que seja
    idêntica ao valor do produto menos desconto, como no C175 analítico.
    """
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        abas = [nome for nome in livro.sheetnames if _slug(nome).startswith("c170")]
        if not abas:
            return pd.DataFrame(columns=COLUNAS_C175)
        regimes = _regimes_0110(livro)
        from auditor_sped_integrado.prazo import classificar_competencia
        mapa_cfop = Counter()
        cobertura_mensal = Counter()
        recorte = Counter()
        grupos = defaultdict(lambda: {
            "registros": 0, "operacao": Decimal(0), "desconto": Decimal(0),
            "base_pis": Decimal(0), "pis": Decimal(0), "base_cofins": Decimal(0),
            "cofins": Decimal(0), "conferidos": 0, "divergentes": 0,
            "chaves_invalidas": 0, "duplicados": 0, "base_pis_ok": Decimal(0),
            "base_cofins_ok": Decimal(0), "pis_ok": Decimal(0),
            "cofins_ok": Decimal(0), "primeira_aba": None, "primeira_linha": None,
        })
        vistos: dict[tuple[str, str, str, str], tuple] = {}
        for nome in abas:
            aba = livro[nome]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            h = next(linhas)
            mes_i, ano_i = _indice(h, "Mês"), _indice(h, "Ano")
            cnpj_i = _indice(h, "CNPJ Estabelecimento(C010)", "CNPJ Estabelecimento", "CNPJ")
            op_i = _indice(h, "Indicador de Operação(C100)", "Indicador de Operação")
            sit_i = _indice(h, "Situação(C100)", "Situação")
            cfop_i = _indice(h, "CFOP")
            cst_pis_i, cst_cof_i = _indice(h, "CST de Pis"), _indice(h, "CST de Cofins")
            key_i, item_i = _indice(h, "Chave(C100)", "Chave"), _indice(h, "Numeração Sequencial", "Item")
            produto_i, desc_i = _indice(h, "Valor Total do Produto"), _indice(h, "Valor de Desconto")
            bp_i, pp_i, vp_i = _indice(h, "Base de Pis"), _indice(h, "Alíquota de Pis"), _indice(h, "Valor de Pis")
            bc_i, pc_i, vc_i = _indice(h, "Base de Cofins"), _indice(h, "Alíquota de Cofins"), _indice(h, "Valor de Cofins")
            for numero_linha, linha in enumerate(linhas, 2):
                periodo = _competencia(linha[mes_i], linha[ano_i])
                if data_referencia is not None:
                    estado_prazo = classificar_competencia(periodo, data_referencia)
                    recorte[estado_prazo] += 1
                    if estado_prazo != "dentro":
                        continue
                cfop = re.sub(r"\D", "", str(linha[cfop_i] or ""))
                if (len(cfop) == 4 and cfop[0] in {"5", "6"}
                        and str(linha[op_i] or "").strip().removesuffix(".0") == "1"
                        and str(linha[sit_i] or "").strip().removesuffix(".0").zfill(2) in {"00", "01"}):
                    mapa_cfop[("C170", cfop, str(linha[cst_pis_i] or "").strip().removesuffix(".0").zfill(2),
                               str(linha[cst_cof_i] or "").strip().removesuffix(".0").zfill(2))] += 1
                    cobertura_mensal[periodo] += 1
                if cfop != "5405" or str(linha[op_i] or "").strip().removesuffix(".0") != "1":
                    continue
                if str(linha[sit_i] or "").strip().removesuffix(".0").zfill(2) not in {"00", "01"}:
                    continue
                cst_pis = str(linha[cst_pis_i] or "").strip().removesuffix(".0").zfill(2)
                cst_cof = str(linha[cst_cof_i] or "").strip().removesuffix(".0").zfill(2)
                if cst_pis != "01" or cst_cof != "01":
                    continue
                periodo = _competencia(linha[mes_i], linha[ano_i])
                cnpj = _cnpj(linha[cnpj_i])
                if not periodo or not cnpj:
                    raise ValueError(f"Identificação incompleta em {nome}, linha {numero_linha}.")
                chave = re.sub(r"\D", "", str(linha[key_i] or ""))
                item = str(linha[item_i] or "").strip().removesuffix(".0")
                uf = chave[:2] if len(chave) == 44 else ""
                chave_ok = (len(chave) == 44 and uf in CODIGOS_UF and chave[6:20] == cnpj
                            and chave[2:6] == periodo[2:4] + periodo[5:7] and item.isdigit() and int(item) > 0)
                pis_presente = linha[pp_i] is not None and str(linha[pp_i]).strip() != ""
                cof_presente = linha[pc_i] is not None and str(linha[pc_i]).strip() != ""
                ap = _decimal(linha[pp_i])
                ac = _decimal(linha[pc_i])
                escolha = (aliquotas_informadas or {}).get(
                    (periodo, cnpj, float(ap) if pis_presente else None,
                     float(ac) if cof_presente else None, "C170"))
                if escolha is not None:
                    if not pis_presente:
                        ap = Decimal(str(escolha[0]))
                    if not cof_presente:
                        ac = Decimal(str(escolha[1]))
                fp = "C170" if pis_presente else "Informada pelo usuário" if ap > 0 else "Ausente"
                fc = "C170" if cof_presente else "Informada pelo usuário" if ac > 0 else "Ausente"
                grupo = grupos[(periodo, cnpj, uf, chave_ok, regimes.get(periodo, ""),
                                 cfop, cst_pis, cst_cof, ap, ac, fp, fc)]
                identidade = (periodo, cnpj, chave, item)
                if chave_ok and identidade in vistos:
                    grupo["duplicados"] += 1
                    primeiro_grupo = grupos[vistos[identidade]]
                    if primeiro_grupo is not grupo:
                        primeiro_grupo["duplicados"] += 1
                    if coletor_nfe is not None:
                        coletor_nfe.adicionar(
                            competencia=periodo, cnpj=cnpj, chave=chave, chave_ok=chave_ok,
                            uf=uf, regime=regimes.get(periodo, ""), aliquota_pis=ap,
                            aliquota_cofins=ac, operacao=Decimal(0), desconto=Decimal(0),
                            base_pis=Decimal(0), pis=Decimal(0), base_cofins=Decimal(0),
                            cofins=Decimal(0), base_ok=True, duplicado=True, aba=nome,
                            linha=numero_linha, fonte_pis=fp, fonte_cofins=fc,
                            fonte_registro="C170",
                        )
                    continue
                if chave_ok:
                    vistos[identidade] = (periodo, cnpj, uf, chave_ok, regimes.get(periodo, ""),
                                         cfop, cst_pis, cst_cof, ap, ac, fp, fc)
                operacao = _decimal(linha[produto_i])
                desconto = _decimal(linha[desc_i])
                bp, bc = _decimal(linha[bp_i]), _decimal(linha[bc_i])
                vp, vc = _decimal(linha[vp_i]), _decimal(linha[vc_i])
                base_ok = (
                    bp > 0 and bc > 0 and abs(bp - bc) <= Decimal("0.01")
                    and vp > 0 and vc > 0
                    and abs(bp * ap / Decimal(100) - vp) <= Decimal("0.02")
                    and abs(bc * ac / Decimal(100) - vc) <= Decimal("0.02")
                )
                grupo["registros"] += 1
                for campo, valor in (("operacao", operacao), ("desconto", desconto),
                                     ("base_pis", bp), ("pis", vp), ("base_cofins", bc), ("cofins", vc)):
                    grupo[campo] += valor
                if base_ok:
                    grupo["conferidos"] += 1
                    grupo["base_pis_ok"] += bp
                    grupo["base_cofins_ok"] += bc
                    grupo["pis_ok"] += vp
                    grupo["cofins_ok"] += vc
                else:
                    grupo["divergentes"] += 1
                grupo["chaves_invalidas"] += int(not chave_ok)
                if coletor_nfe is not None:
                    coletor_nfe.adicionar(
                        competencia=periodo, cnpj=cnpj, chave=chave, chave_ok=chave_ok,
                        uf=uf, regime=regimes.get(periodo, ""), aliquota_pis=ap,
                        aliquota_cofins=ac, operacao=operacao, desconto=desconto,
                        base_pis=bp, pis=vp, base_cofins=bc, cofins=vc,
                        base_ok=base_ok, duplicado=False, aba=nome, linha=numero_linha,
                        fonte_pis=fp, fonte_cofins=fc, fonte_registro="C170",
                    )
                if grupo["primeira_linha"] is None:
                    grupo["primeira_aba"], grupo["primeira_linha"] = nome, numero_linha
        saida = []
        for (periodo, cnpj, uf, chave_ok, regime, cfop, cst_pis, cst_cof, ap, ac, fp, fc), d in grupos.items():
            saida.append({
                "competencia": periodo, "cnpj_estabelecimento": cnpj, "uf_nfe": uf,
                "chave_conferida": chave_ok, "regime_0110": regime, "cfop": cfop,
                "cst_pis": cst_pis, "cst_cofins": cst_cof,
                "aliquota_pis": float(ap), "aliquota_cofins": float(ac),
                "fonte_aliquota_pis": fp, "fonte_aliquota_cofins": fc,
                "registros_c175": d["registros"], "valor_operacao": float(d["operacao"]),
                "valor_desconto": float(d["desconto"]), "base_pis": float(d["base_pis"]),
                "valor_pis": float(d["pis"]), "base_cofins": float(d["base_cofins"]),
                "valor_cofins": float(d["cofins"]), "registros_base_conferida": d["conferidos"],
                "base_pis_conferida": float(d["base_pis_ok"]),
                "base_cofins_conferida": float(d["base_cofins_ok"]),
                "valor_pis_conferido": float(d["pis_ok"]),
                "valor_cofins_conferido": float(d["cofins_ok"]),
                "registros_base_divergente": d["divergentes"],
                "registros_chave_invalida": d["chaves_invalidas"],
                "registros_duplicados_entre_abas": d["duplicados"],
                "primeira_aba": d["primeira_aba"], "primeira_linha_c175": d["primeira_linha"],
                "fonte_registro": "C170",
            })
        resultado = pd.DataFrame(saida, columns=COLUNAS_C175)
        resultado.attrs["mapa_cfop"] = mapa_cfop
        resultado.attrs["cobertura_mensal"] = dict(cobertura_mensal)
        resultado.attrs["recorte_prazo"] = dict(recorte)
        return resultado
    finally:
        livro.close()


def diagnosticar_st_c190(arquivo: str | Path | bytes) -> pd.DataFrame:
    """Verifica ST no agregado C190 do ICMS/IPI, sem somar com C100/C170."""
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        abas_c190 = [titulo for titulo in livro.sheetnames if _slug(titulo).startswith("c190")]
        if not abas_c190:
            raise ValueError("Aba C190 não encontrada no ICMS/IPI.")
        grupos = defaultdict(lambda: {"registros": 0, "valor": Decimal("0"), "primeira_linha": None})
        for nome in abas_c190:
            aba = livro[nome]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            cabecalho = next(linhas)
            mes_i, ano_i = _indice(cabecalho, "Mês"), _indice(cabecalho, "Ano")
            cnpj_i = _indice(cabecalho, "CNPJ")
            oper_i = _indice(cabecalho, "Indicador de Operação(C100)")
            cfop_i = _indice(cabecalho, "CFOP")
            try:
                st_i = _indice(cabecalho, "Valor de Icms ST", "Valor de ICMS - ST")
            except ValueError:
                continue
            for numero_linha, linha in enumerate(linhas, start=2):
                valor_st = _decimal(linha[st_i])
                if valor_st <= 0:
                    continue
                periodo = _competencia(linha[mes_i], linha[ano_i])
                cnpj = _cnpj(linha[cnpj_i])
                operacao = str(linha[oper_i] if linha[oper_i] is not None else "").strip().removesuffix(".0")
                cfop = re.sub(r"\D", "", str(linha[cfop_i] or ""))
                if not periodo or not cnpj or operacao not in {"0", "1"} or len(cfop) != 4:
                    raise ValueError(f"Identificação incompleta em {nome}, linha {numero_linha}.")
                grupo = grupos[(periodo, cnpj, cfop, operacao)]
                grupo["registros"] += 1
                grupo["valor"] += valor_st
                if grupo["primeira_linha"] is None:
                    grupo["primeira_linha"] = numero_linha
        saida = [
            {
                "competencia": periodo, "cnpj_estabelecimento": cnpj, "cfop": cfop,
                "indicador_operacao": operacao, "registros_c190": dados["registros"],
                "valor_icms_st_destacado": float(dados["valor"].quantize(Decimal("0.01"))),
                "primeira_linha_c190": dados["primeira_linha"],
            }
            for (periodo, cnpj, cfop, operacao), dados in grupos.items()
        ]
        return pd.DataFrame(saida, columns=COLUNAS_C190).sort_values(
            ["competencia", "cnpj_estabelecimento", "cfop", "indicador_operacao"], ignore_index=True
        ) if saida else pd.DataFrame(columns=COLUNAS_C190)
    finally:
        livro.close()


def mapear_saidas_c190(arquivo: str | Path | bytes, data_referencia=None,
                       cobertura_mensal: Counter | None = None) -> list[dict]:
    """Mostra CFOPs de saída do C190 sem supor CST de PIS/COFINS."""
    fonte = BytesIO(arquivo) if isinstance(arquivo, bytes) else Path(arquivo)
    livro = load_workbook(fonte, read_only=True, data_only=True)
    try:
        from auditor_sped_integrado.prazo import classificar_competencia
        contagens = Counter()
        for nome in (n for n in livro.sheetnames if _slug(n).startswith("c190")):
            aba = livro[nome]
            aba.reset_dimensions()
            linhas = aba.iter_rows(values_only=True)
            h = next(linhas)
            mes_i, ano_i = _indice(h, "Mês"), _indice(h, "Ano")
            op_i = _indice(h, "Indicador de Operação(C100)", "Indicador de Operação")
            cfop_i = _indice(h, "CFOP")
            for linha in linhas:
                if data_referencia is not None and classificar_competencia(
                    _competencia(linha[mes_i], linha[ano_i]), data_referencia
                ) != "dentro":
                    continue
                if str(linha[op_i] or "").strip().removesuffix(".0") != "1":
                    continue
                cfop = re.sub(r"\D", "", str(linha[cfop_i] or ""))
                if len(cfop) == 4 and cfop[0] in {"5", "6"}:
                    contagens[cfop] += 1
                    if cobertura_mensal is not None:
                        cobertura_mensal[_competencia(linha[mes_i], linha[ano_i])] += 1
        return [
            {"Registro": "C190 ICMS/IPI", "CFOP": cfop, "CST PIS": "Não consta",
             "CST COFINS": "Não consta", "Linhas": total, "No cálculo 1125": False}
            for cfop, total in sorted(contagens.items())
        ]
    finally:
        livro.close()
