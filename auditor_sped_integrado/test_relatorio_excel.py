from datetime import date, datetime
from io import BytesIO
import unittest

import pandas as pd
from openpyxl import load_workbook

from auditor_sped_integrado.conferencia_nfe import ConferenciaNFe
from auditor_sped_integrado.relatorio_excel import gerar_relatorio_excel, resumir_candidatos


class RelatorioExcelTest(unittest.TestCase):
    def test_relatorio_separa_contribuicao_escriturada_de_credito(self):
        vendas = pd.DataFrame([
        {
            "competencia": "2025-01", "cnpj_estabelecimento": "10693312000505",
            "cfop": "5405", "cst_pis": "01", "cst_cofins": "01",
            "aliquota_pis": 1.65, "aliquota_cofins": 7.6, "registros_c175": 2,
            "valor_operacao": 100.0, "base_pis": 80.0, "valor_pis": 1.32,
            "base_cofins": 80.0, "valor_cofins": 6.08,
            "primeira_aba": "C175", "primeira_linha_c175": 2,
        },
        {
            "competencia": "2025-01", "cnpj_estabelecimento": "10693312000505",
            "cfop": "5405", "cst_pis": "04", "cst_cofins": "04",
            "aliquota_pis": 0.0, "aliquota_cofins": 0.0, "registros_c175": 1,
            "valor_operacao": 50.0, "base_pis": 0.0, "valor_pis": 0.0,
            "base_cofins": 0.0, "valor_cofins": 0.0,
            "primeira_aba": "C175", "primeira_linha_c175": 3,
        },
        ])
        candidatos = resumir_candidatos(vendas)
        self.assertEqual(len(candidatos), 1)
        self.assertEqual(candidatos.iloc[0]["pis_escriturado"], 1.32)
        self.assertEqual(candidatos.iloc[0]["contribuicoes_sob_analise"], 7.4)
        self.assertIsNone(candidatos.iloc[0]["credito_pis_calculado"])
        vazio = pd.DataFrame()
        from decimal import Decimal
        with ConferenciaNFe() as conferencia:
            for valor in (Decimal("40"), Decimal("40")):
                conferencia.adicionar(
                    competencia="2025-01", cnpj="10693312000505",
                    chave="51250110693312000505550010000000011000000000", chave_ok=True,
                    uf="51", regime="1", aliquota_pis=Decimal("1.65"),
                    aliquota_cofins=Decimal("7.6"), operacao=valor, desconto=Decimal("0"),
                    base_pis=valor, pis=Decimal("0.66"), base_cofins=valor,
                    cofins=Decimal("3.04"), base_ok=True, duplicado=False,
                    aba="C175", linha=2,
                )
            dados = gerar_relatorio_excel(
                [{"tipo": "EFD-Contribuições", "arquivo": "origem.xlsx", "sha256": "abc",
                  "abas": {"C175": 2}, "linhas": 2}],
                vazio, vazio, vazio, vendas, vazio, "Lucro real", datetime(2026, 10, 21, 15, 0),
                conferencia_nfe=conferencia,
            )
        livro = load_workbook(BytesIO(dados), read_only=False, data_only=False)
        self.assertEqual(livro.sheetnames, ["Resumo_Executivo", "Memoria_Mensal", "Conferencia_NF"])
        resumo = livro["Resumo_Executivo"]
        linha_total = next(linha for linha in resumo if linha[0].value == "Total")
        self.assertEqual(linha_total[1].value, 1.32)
        self.assertEqual(linha_total[2].value, 6.08)
        linha_documentada = next(linha for linha in resumo if linha[0].value == "Após conferência manual")
        self.assertEqual(linha_documentada[1].data_type, "f")
        nf = livro["Conferencia_NF"]
        self.assertEqual(nf["AB2"].value, 2)
        self.assertEqual(nf["K2"].value, 80)
        self.assertEqual(nf["F2"].value, "Apto para conferência")
        self.assertEqual(nf["T2"].value, "Pendente")
        self.assertTrue(nf["U2"].value.startswith("=IF("))
        self.assertEqual(nf["S2"].value, "Pendente")
        self.assertIn('S2="Conferido"', nf["U2"].value)
        self.assertIn('MIN(P2,K2)*M2/100', nf["U2"].value)
        self.assertEqual(nf["Y1"].value, "Aba de origem")
        self.assertEqual(livro["Memoria_Mensal"]["G1"].value, "Registros de origem")
        self.assertEqual(nf.freeze_panes, "D2")
        self.assertEqual(nf["O2"].fill.fgColor.rgb[-6:], "FFF2CC")
        self.assertEqual(len(nf.data_validations.dataValidation), 2)
        self.assertEqual(len(nf.conditional_formatting), 2)
        livro.close()

    def test_cobertura_sinaliza_mes_sem_grupo_sem_inventar_valor(self):
        vazio = pd.DataFrame()
        dados = gerar_relatorio_excel(
            [], vazio, vazio, vazio, vazio, vazio, "A confirmar",
            data_referencia=date(2026, 10, 22),
        )
        livro = load_workbook(BytesIO(dados), read_only=True, data_only=True)
        memoria = livro["Memoria_Mensal"]
        meses_sem_grupo = [linha for linha in memoria.values if linha[0] == "2025-01"]
        self.assertEqual(len(meses_sem_grupo), 1)
        self.assertIsNone(meses_sem_grupo[0][13])
        self.assertIn("Sem grupo no filtro", meses_sem_grupo[0][17])
        livro.close()
